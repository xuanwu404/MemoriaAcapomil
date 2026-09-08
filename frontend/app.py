from __future__ import annotations

import io
import importlib
import re
import sqlite3
import shutil
import sys
import unicodedata
import zipfile
from pathlib import Path
from typing import Any

import pandas as pd
import streamlit as st
import streamlit.components.v1 as components
from docx import Document


ROOT_DIR = Path(__file__).resolve().parent.parent
if str(ROOT_DIR) not in sys.path:
	sys.path.append(str(ROOT_DIR))

from backend.database import (
	actualizar_unidad,
	actualizar_usuario,
	buscar_alumno_por_rut_anio_mencion,
	crear_alumno,
	crear_unidad,
	crear_usuario,
	eliminar_alumno,
	eliminar_evaluacion,
	eliminar_retroalimentacion_local,
	eliminar_usuario_por_username,
	eliminar_usuario,
	eliminar_unidad,
	formatear_rut_chileno,
	guardar_evaluacion,
	guardar_mapeo_pregunta,
	guardar_puntajes_pregunta,
	guardar_retroalimentacion_local,
	inicializar_si_no_existe,
	cargar_indicadores_desde_df,
	extraer_indicadores_desde_pdf,
	listar_evaluaciones_unidad,
	listar_alumnos,
	obtener_historial_retroalimentaciones,
	obtener_puntajes_pregunta,
	listar_indicadores_unidad,
	listar_mapeo_evaluacion,
	listar_resultados_alumno,
	listar_unidades,
	obtener_usuarios,
	obtener_alumnos,
	obtener_brechas_alumno,
	obtener_unidad_por_id,
	obtener_usuario_por_credenciales,
	obtener_usuario_por_username,
	reemplazar_indicadores_unidad,
	reemplazar_mapeo_evaluacion,
	reemplazar_resultados_alumno,
	validar_rut_chileno,
)
from backend.orquestador import (
	generar_plan_trabajo_alumno,
	generar_plan_trabajo_cohorte,
	generar_sugerencia_pregunta_cobertura,
)
from backend.pipeline_rag import cargar_documentos_docx, construir_vector_store


DOCENTE_DEFAULT = "TCL Pablo Varas Briones"
ANIOS_CRIM = ["I CRIM", "II CRIM", "III CRIM", "IV CRIM"]
MENCIONES = [
	"Comunicaciones e Informatica",
	"Sistemas de Armas",
	"Abastecimiento y Finanzas",
	"Geografia",
]


st.set_page_config(
	page_title="ACAPOMIL | Contextos de Asignatura",
	page_icon="🎯",
	layout="wide",
)


def _bloquear_lector_navegador() -> None:
	"""Inyecta un listener JS en el documento padre para evitar que el navegador abra PDFs sueltos por error."""
	components.html(
		"""
		<script>
		const doc = window.parent.document;

		doc.addEventListener('dragover', function(e) {
			e.preventDefault();
		}, false);

		doc.addEventListener('drop', function(e) {
			// Verifica si el objetivo del drop es la zona de carga de Streamlit
			const isDropzone = e.target.closest('[data-testid="stFileUploadDropzone"]');

			// Si cayó fuera del área de carga, bloqueamos la acción nativa del navegador
			if (!isDropzone) {
				e.preventDefault();
				e.stopPropagation();
			}
		}, false);
		</script>
		""",
		height=0,
		width=0,
	)


def _inject_styles() -> None:
	st.markdown(
		"""
		<style>
		/* 1. Fondo global y texto base */
		.stApp { background-color: #F8F9FA; color: #212529; }

		/* 2. Cabeceras Hero (Rojo Institucional con texto 100% BLANCO) */
		.hero {
			background: linear-gradient(135deg, #B71C1C 0%, #D32F2F 100%);
			padding: 1.5rem; border-radius: 8px;
			box-shadow: 0 4px 6px rgba(0,0,0,0.1);
			margin-bottom: 2rem; border: 1px solid #880E4F;
		}
		.hero h2, .hero p, .hero h3 { color: #FFFFFF !important; margin: 0; text-shadow: 1px 1px 2px rgba(0,0,0,0.3); }

		/* 3. Cajas y Contenedores (Bordes definidos y sombras sutiles) */
		div[data-testid="stVerticalBlock"] > div[style*="border"] {
			background-color: #FFFFFF !important;
			border-radius: 8px !important;
			border: 1px solid #CED4DA !important; /* Borde claramente visible */
			box-shadow: 0 2px 4px rgba(0,0,0,0.05) !important;
			padding: 1.5rem !important;
			margin-bottom: 1.5rem !important;
		}

		/* 4. Inputs y Selectboxes (Affordance: deben parecer clickeables) */
		.stTextInput input, .stSelectbox div[data-baseweb="select"], .stTextArea textarea {
			border: 1px solid #ADB5BD !important;
			border-radius: 6px !important;
			background-color: #FFFFFF !important;
			color: #212529 !important;
		}
		.stTextInput input:focus, .stSelectbox div[data-baseweb="select"]:focus-within, .stTextArea textarea:focus {
			border-color: #D32F2F !important;
			box-shadow: 0 0 0 2px rgba(211, 47, 47, 0.25) !important;
		}

		/* 5. Botones Primarios */
		button[kind="primary"] {
			background-color: #D32F2F !important;
			color: #FFFFFF !important;
			border: none !important; border-radius: 6px !important;
			font-weight: 600 !important; padding: 0.5rem 1rem !important;
		}
		button[kind="primary"]:hover { background-color: #B71C1C !important; box-shadow: 0 4px 8px rgba(183,28,28,0.3) !important;}

		/* 6. Sidebar (Navegación limpia y botones uniformes) */
		section[data-testid="stSidebar"] { background-color: #F8F9FA !important; border-right: 1px solid #DEE2E6 !important; }

		/* Ocultar el círculo nativo del radio button */
		section[data-testid="stSidebar"] .stRadio div[role="radio"] div:first-child,
		section[data-testid="stSidebar"] .stRadio input[type="radio"] { display: none !important; }

		/* Contenedor del grupo de botones */
		section[data-testid="stSidebar"] .stRadio > div { width: 100% !important; gap: 8px !important; }

		/* Diseño de cada botón (Ancho total uniforme) */
		section[data-testid="stSidebar"] .stRadio label[data-baseweb="radio"] {
			display: flex !important;
			width: 100% !important;
			background-color: #FFFFFF !important;
			border: 1px solid #CED4DA !important;
			border-radius: 6px !important;
			padding: 12px 15px !important;
			margin: 0 !important;
			transition: all 0.2s ease-in-out !important;
			box-sizing: border-box !important;
		}
		section[data-testid="stSidebar"] .stRadio label[data-baseweb="radio"] p {
			color: #495057 !important; /* Gris oscuro legible */
			font-weight: 500 !important;
			margin: 0 !important;
			width: 100% !important;
		}
		section[data-testid="stSidebar"] .stRadio label[data-baseweb="radio"]:hover { background-color: #E9ECEF !important; }

		/* 7. Sidebar Activo (Rojo con texto BLANCO OBLIGATORIO) */
		section[data-testid="stSidebar"] .stRadio label[data-baseweb="radio"]:has(input[type="radio"]:checked) {
			background-color: #D32F2F !important; border-color: #B71C1C !important;
		}
		section[data-testid="stSidebar"] .stRadio label[data-baseweb="radio"]:has(input[type="radio"]:checked) * {
			color: #FFFFFF !important; font-weight: 600 !important;
		}

		/* 8. Etiquetas / Tags */
		.tag { background: #FFEBEE; border: 1px solid #FFCDD2; color: #C62828; padding: 0.2rem 0.8rem; border-radius: 12px; font-weight: 600; font-size: 0.85rem; display: inline-block; margin: 0.2rem; }

		/* 9. File Uploader (Drag & Drop Dropzone Mejorado) */
		[data-testid="stFileUploadDropzone"] {
			border: 3px dashed #B71C1C !important;
			background-color: #F8F9FA !important;
			border-radius: 12px !important;
			padding: 3.5rem 2rem !important; /* Hace la caja considerablemente más alta */
			transition: all 0.3s ease-in-out !important;
		}

		/* Efecto Hover / Drag Over: Crece y cambia de color para atraer el archivo */
		[data-testid="stFileUploadDropzone"]:hover,
		[data-testid="stFileUploadDropzone"]:focus-within {
			background-color: #FFEBEE !important;
			border-color: #D32F2F !important;
			transform: scale(1.02) !important; /* Aumenta su tamaño un 2% al acercar el archivo */
			box-shadow: 0 8px 16px rgba(211, 47, 47, 0.15) !important;
		}

		/* Estilizar el texto interior y el icono para que hagan juego */
		[data-testid="stFileUploadDropzone"] div[data-testid="stMarkdownContainer"] p {
			color: #333333 !important;
			font-size: 1.1rem !important;
			font-weight: 500 !important;
		}

		/* Estilizar el botón interno "Browse files" para que sea coherente */
		[data-testid="stFileUploadDropzone"] button {
			background-color: #D32F2F !important;
			color: #FFFFFF !important;
			border: none !important;
			border-radius: 6px !important;
			font-weight: 600 !important;
			transition: all 0.2s ease !important;
		}
		[data-testid="stFileUploadDropzone"] button:hover {
			background-color: #B71C1C !important;
			transform: translateY(-1px) !important;
			box-shadow: 0 4px 8px rgba(183, 28, 28, 0.3) !important;
		}
		</style>
		""",
		unsafe_allow_html=True,
	)


def _init_state() -> None:
	inicializar_si_no_existe()
	if "is_logged" not in st.session_state:
		st.session_state.is_logged = False
	if "usuario_activo" not in st.session_state:
		st.session_state.usuario_activo = ""
	if "username_activo" not in st.session_state:
		st.session_state.username_activo = ""
	if "docente_id" not in st.session_state:
		st.session_state.docente_id = None
	if "usuario_rol" not in st.session_state:
		st.session_state.usuario_rol = ""
	if "unidad_activa_id" not in st.session_state:
		st.session_state.unidad_activa_id = None
	if "borrador_retro" not in st.session_state:
		st.session_state.borrador_retro = ""
	if "evaluacion_activa_id" not in st.session_state:
		st.session_state.evaluacion_activa_id = None
	if "uploader_key" not in st.session_state:
		st.session_state.uploader_key = 1
	if "indicadores_excel_evento" not in st.session_state:
		st.session_state.indicadores_excel_evento = ""
	if "indicadores_excel_procesado" not in st.session_state:
		st.session_state.indicadores_excel_procesado = False
	if "wizard_eval_step" not in st.session_state:
		st.session_state.wizard_eval_step = 1


def _slug(texto: str) -> str:
	return "".join(c.lower() if c.isalnum() else "_" for c in texto).strip("_")


def _data_unidad_dir(unidad_id: str) -> Path:
	return ROOT_DIR / "data" / "doctrina" / unidad_id


def _vector_unidad_dir(unidad_id: str) -> Path:
	return ROOT_DIR / "data" / "vector_store" / unidad_id


def _eval_unidad_dir(unidad_id: str) -> Path:
	return ROOT_DIR / "data" / "evaluaciones" / unidad_id


def _limpiar_texto_pdf(texto: str) -> str:
	if not texto:
		return ""
	return " ".join(str(texto).replace("\x00", "").split())


def _leer_texto_docx_desde_bytes(archivo_bytes: bytes) -> str:
	partes: list[str] = []
	try:
		documento = Document(io.BytesIO(archivo_bytes))
		for parrafo in documento.paragraphs:
			texto_parrafo = str(parrafo.text or "").strip()
			if texto_parrafo and not _es_linea_portada_instrumento(texto_parrafo):
				partes.append(texto_parrafo)
	except Exception as e:
		print(f"[DOCX] Error al extraer texto: {e}")
		return "\n".join(partes).strip()
	return "\n".join(partes).strip()


def _leer_texto_pdf_desde_bytes(archivo_bytes: bytes) -> str:
	try:
		pypdf_module = importlib.import_module("pypdf")
		reader = pypdf_module.PdfReader(io.BytesIO(archivo_bytes))
		partes: list[str] = []
		for pagina in reader.pages[:3]:
			partes.append(pagina.extract_text() or "")
		return "\n".join(partes).strip()
	except Exception:
		return ""


def _normalizar_texto_parser(texto: Any) -> str:
	valor = str(texto or "").strip().lower()
	valor = "".join(
		c for c in unicodedata.normalize("NFKD", valor) if not unicodedata.combining(c)
	)
	valor = re.sub(r"\s+", " ", valor)
	return valor


def _clave_orden_pregunta(pregunta_str: str) -> tuple[int, int]:
	roman_map = {
		"I": 1,
		"II": 2,
		"III": 3,
		"IV": 4,
		"V": 5,
		"VI": 6,
		"VII": 7,
		"VIII": 8,
		"IX": 9,
		"X": 10,
	}
	partes = str(pregunta_str).strip().upper().split(".")
	if len(partes) == 2:
		roman, num_str = partes
		val_romano = roman_map.get(roman, 99)
		try:
			val_num = int(num_str)
		except ValueError:
			val_num = 0
		return (val_romano, val_num)
	return (99, 0)


def _normalizar_codigo_req(codigo: str) -> str:
	texto = str(codigo or "").upper()
	texto = texto.replace(",", ".")
	texto = re.sub(r"[^A-Z0-9.]+", "", texto)
	return texto


def _es_req_valido(req: str) -> bool:
	norm = _normalizar_codigo_req(req)
	return bool(re.match(r"^R\d+(?:\.\d+)*$", norm))


def _extraer_preguntas_limpias(texto_celda: str) -> list[str]:
	texto = str(texto_celda or "").upper().replace("\n", " ")
	partes = texto.split(",")
	preguntas: list[str] = []
	vistas: set[str] = set()
	for parte in partes:
		fragmento = parte.strip()
		if not fragmento:
			continue
		match = re.search(r"\b([IVX]+|\d{1,2})\.(\d{1,2})\b", fragmento)
		if not match:
			continue
		izquierda, derecha = match.group(1), match.group(2)
		if int(derecha) <= 0:
			continue
		pregunta = f"{izquierda}.{int(derecha)}"
		if pregunta in vistas:
			continue
		vistas.add(pregunta)
		preguntas.append(pregunta)
	return preguntas


def _es_linea_portada_instrumento(linea: str) -> bool:
	normalizada = _normalizar_texto_parser(linea)
	if not normalizada:
		return True
	frases_bloqueadas = [
		"req",
		"criterios de evaluacion",
		"indicador de evaluacion",
		"aprendizaje esperado por evaluar",
		"tareas asociadas",
		"requerimientos",
		"indicadores de desempeno",
		"unidad de aprendizaje",
		"ejercito de chile",
		"division educacion",
	]
	return any(frase in normalizada for frase in frases_bloqueadas)


def _limpiar_enunciado_pregunta(texto: str) -> str:
	texto_limpio = _limpiar_texto_pdf(texto)
	texto_limpio = texto_limpio.lstrip(". ")
	texto_limpio = re.sub(r"^[\s\.-:;•·,]+", "", texto_limpio)
	texto_limpio = re.sub(r"\s+", " ", texto_limpio).strip()
	return texto_limpio


def _extraer_preguntas_docx_estructuradas(archivo_bytes: bytes) -> list[dict[str, str]]:
	try:
		documento = Document(io.BytesIO(archivo_bytes))
	except Exception as e:
		print(f"[DOCX] Error al abrir el documento: {e}")
		return []

	parrafos: list[str] = []
	for parrafo in documento.paragraphs:
		texto = _limpiar_enunciado_pregunta(parrafo.text)
		if not texto:
			continue
		if _es_linea_portada_instrumento(texto):
			continue
		if re.fullmatch(r"[_\-]{4,}", texto):
			continue
		parrafos.append(texto)

	if not parrafos:
		return []

	patron_item = re.compile(r"^ITEM\s+([IVX]+)\b", re.IGNORECASE)
	patron_pregunta = re.compile(r"^(?:\d+\.\s*)?(¿|[A-Za-zÁÉÍÓÚÑáéíóúñ].*?\?)")
	patron_opcion = re.compile(r"^\s*([a-eA-E])[\)\.\-]\s*(.+)$")
	patron_vf = re.compile(r"^\s*(?:_+\s*)?(.+)$")
	patron_solo_subrayado = re.compile(r"^[_\-\s]+$")
	patron_fin_iii = re.compile(r"^(?:PREGUNTA\s*R\s*4\.1|PREGUNTA\s*R\s*\d|PONDERACI[ÓO]N|NOTA\s+FINAL|PUNTAJE)", re.IGNORECASE)
	patron_inicio_requerimiento = re.compile(r"^\s*(?:\d+\.|[-*•])\s*(.+)$")

	registros: list[dict[str, str]] = []
	seccion_actual: str | None = None
	contador_i = 1
	contador_ii = 1
	contador_iii = 1
	buffer_pregunta_i: list[str] = []
	buffer_opciones_i: list[str] = []
	buffer_iii_contexto: list[str] = []
	buffer_iii_requerimientos: list[str] = []
	capturando_iii = False
	modo_iii_ejercicio = False

	def _cerrar_item_i() -> None:
		nonlocal buffer_pregunta_i, buffer_opciones_i, contador_i
		texto = _limpiar_enunciado_pregunta(" ".join(buffer_pregunta_i + buffer_opciones_i))
		if texto:
			registros.append({"codigo": f"I.{contador_i}", "enunciado": texto, "display": f"Ítem I.{contador_i}: {texto}"})
			contador_i += 1
		buffer_pregunta_i = []
		buffer_opciones_i = []

	def _cerrar_item_ii(texto: str) -> None:
		nonlocal contador_ii
		texto_limpio = _limpiar_enunciado_pregunta(texto)
		if texto_limpio:
			registros.append({"codigo": f"II.{contador_ii}", "enunciado": texto_limpio, "display": f"Ítem II.{contador_ii}: {texto_limpio}"})
			contador_ii += 1

	def _cerrar_item_iii(texto: str) -> None:
		nonlocal contador_iii
		texto_limpio = _limpiar_enunciado_pregunta(texto)
		if texto_limpio:
			registros.append({"codigo": f"III.{contador_iii}", "enunciado": texto_limpio, "display": f"Ítem III.{contador_iii}: {texto_limpio}"})
			contador_iii += 1

	def _agregar_parrafo_fallback(texto: str) -> None:
		nonlocal contador_iii
		texto_limpio = _limpiar_enunciado_pregunta(texto)
		if texto_limpio:
			registros.append({"codigo": f"P{contador_iii}", "enunciado": texto_limpio, "display": f"Ítem P{contador_iii}: {texto_limpio}"})
			contador_iii += 1

	for linea in parrafos:
		if patron_solo_subrayado.match(linea):
			continue

		match_item = patron_item.match(linea)
		if match_item:
			if seccion_actual == "I":
				_cerrar_item_i()
			elif seccion_actual == "III":
				if buffer_iii_contexto:
					_cerrar_item_iii(" ".join(buffer_iii_contexto))
				buffer_iii_contexto = []
				buffer_iii_requerimientos = []
				capturando_iii = False
				modo_iii_ejercicio = False
			seccion_actual = str(match_item.group(1)).upper()
			continue

		if seccion_actual == "I":
			match_opcion = patron_opcion.match(linea)
			if match_opcion and buffer_pregunta_i:
				opcion_texto = f"{match_opcion.group(1).lower()}) {match_opcion.group(2).strip()}"
				buffer_opciones_i.append(opcion_texto)
				continue

			if "¿" in linea or patron_pregunta.match(linea):
				if buffer_pregunta_i:
					_cerrar_item_i()
					buffer_pregunta_i = []
					buffer_opciones_i = []
				pregunta_base = re.sub(r"^\d+\.\s*", "", linea).strip()
				buffer_pregunta_i.append(pregunta_base)
				continue

			if buffer_pregunta_i:
				buffer_pregunta_i.append(linea)
				continue

		if seccion_actual == "II":
			if patron_vf.match(linea) or patron_solo_subrayado.match(linea):
				continue
			if re.search(r"puntaje|instruccion|marque|respuesta", _normalizar_texto_parser(linea)):
				continue
			texto_vf = re.sub(r"^_+\s*", "", linea).lstrip(". ")
			if not texto_vf:
				continue
			_cerrar_item_ii(texto_vf)
			continue

		if seccion_actual == "III":
			if patron_fin_iii.match(linea):
				if buffer_iii_contexto:
					_cerrar_item_iii(" ".join(buffer_iii_contexto))
				break
			texto_norm = _limpiar_enunciado_pregunta(linea)
			if not texto_norm:
				continue
			if "PLANTEAMIENTO" in _normalizar_texto_parser(texto_norm):
				capturando_iii = True
				texto_post = texto_norm.split(":", 1)[1].strip() if ":" in texto_norm else ""
				if texto_post:
					buffer_iii_contexto.append(texto_post)
				continue
			if "EJERCICIO" in _normalizar_texto_parser(texto_norm):
				modo_iii_ejercicio = True
				texto_post = texto_norm.split(":", 1)[1].strip() if ":" in texto_norm else ""
				if texto_post:
					buffer_iii_requerimientos.append(texto_post)
				continue
			if not capturando_iii:
				continue
			if modo_iii_ejercicio:
				match_req = patron_inicio_requerimiento.match(texto_norm)
				if match_req:
					buffer_iii_requerimientos.append(match_req.group(1).strip())
				else:
					buffer_iii_requerimientos.append(texto_norm)
			else:
				buffer_iii_contexto.append(texto_norm)
			continue

	if seccion_actual == "I":
		_cerrar_item_i()

	if seccion_actual == "III":
		if buffer_iii_contexto:
			_cerrar_item_iii(" ".join(buffer_iii_contexto))
		for req in buffer_iii_requerimientos:
			_cerrar_item_iii(req)

	if not registros:
		parrafos_limpios = [
			texto
			for texto in parrafos
			if texto and not _es_linea_portada_instrumento(texto) and not re.fullmatch(r"[_\-\s]+", texto)
		]
		for texto in parrafos_limpios:
			_agregar_parrafo_fallback(texto)

	registros_unicos: list[dict[str, str]] = []
	vistos: set[str] = set()
	for registro in registros:
		codigo = str(registro.get("codigo", "")).strip().upper()
		if not codigo or codigo in vistos:
			continue
		vistos.add(codigo)
		registros_unicos.append(registro)

	return registros_unicos


def _leer_texto_instrumento_desde_path(archivo_path: Path) -> str:
	if not archivo_path.exists():
		return ""
	contenido = archivo_path.read_bytes()
	sufijo = archivo_path.suffix.lower()
	if sufijo == ".docx":
		return _leer_texto_docx_desde_bytes(contenido)
	if sufijo == ".pdf":
		return _leer_texto_pdf_desde_bytes(contenido)
	try:
		return contenido.decode("utf-8")
	except UnicodeDecodeError:
		return contenido.decode("latin-1", errors="ignore")


def _detectar_preguntas_instrumento_detalladas(
	texto: str,
	permitir_secuencial: bool = False,
) -> list[dict[str, str]]:
	lineas_crudas = [re.sub(r"\s+", " ", linea.strip()) for linea in str(texto or "").splitlines()]
	lineas = [linea for linea in lineas_crudas if linea and not _es_linea_portada_instrumento(linea)]
	if not lineas:
		return []

	patron_item = re.compile(r"^\s*ITEM\s+([IVX]+)\b", re.IGNORECASE)
	patron_numero = re.compile(r"^\s*(\d{1,2})\.\s*(.*)$")
	patron_opcion = re.compile(r"^\s*([a-eA-E])[\)\.\-]\s*(.*)$")
	patron_vf = re.compile(r"^\s*(?:_+|V\/?F|F\/?V|VERDADERO\s*\/\s*FALSO|V\s*\/\s*F|F\s*\/\s*V)\b\s*(.*)$", re.IGNORECASE)
	patron_bloque_subrayado = re.compile(r"^[_\-\s]{4,}$")
	patron_inicio_enunciado = re.compile(r"(?:^|\s)(?:¿|[A-ZÁÉÍÓÚÑ].*\?)")
	patron_texto_interrogativo = re.compile(r"^(cual|cuál|como|cómo|que|qué|donde|dónde|cuando|cuándo|quien|quién|por que|por qué)\b", re.IGNORECASE)
	patron_respuesta = re.compile(r"^\s*(?:[vVfF]|verdadero|falso|si|sí|no)\s*$", re.IGNORECASE)

	def _es_basura(linea: str) -> bool:
		if not linea:
			return True
		if patron_bloque_subrayado.match(linea):
			return True
		normalizada = _normalizar_texto_parser(linea)
		return not any(ch.isalnum() for ch in normalizada)

	def _es_inicio_pregunta(linea: str) -> bool:
		if patron_numero.match(linea):
			return True
		if patron_vf.match(linea):
			return True
		if "¿" in linea or "?" in linea:
			return True
		if patron_texto_interrogativo.match(linea):
			return True
		if patron_inicio_enunciado.search(linea) and not patron_opcion.match(linea):
			return True
		return False

	def _limpiar_bloque_enunciado(bloque: list[str]) -> str:
		texto_bloque = " ".join(bloque)
		texto_bloque = _limpiar_enunciado_pregunta(texto_bloque)
		texto_bloque = re.sub(r"\s{2,}", " ", texto_bloque).strip()
		return texto_bloque

	def _fallback_por_parrafos(parrafos: list[str]) -> list[dict[str, str]]:
		registros: list[dict[str, str]] = []
		secuencia = 1
		for parrafo in parrafos:
			texto_parrafo = _limpiar_enunciado_pregunta(parrafo)
			if not texto_parrafo or _es_basura(texto_parrafo):
				continue
			codigo = f"P{secuencia}"
			registros.append(
				{
					"codigo": codigo,
					"enunciado": texto_parrafo,
					"display": f"Ítem {codigo}: {texto_parrafo}",
				}
			)
			secuencia += 1
		return registros

	secciones: dict[str, list[str]] = {"I": [], "II": [], "III": []}
	seccion_actual: str | None = None
	for linea in lineas:
		match_item = patron_item.match(linea)
		if match_item:
			seccion_actual = str(match_item.group(1)).upper()
			continue
		if seccion_actual in secciones:
			secciones[seccion_actual].append(linea)

	def _es_item_i(linea: str) -> bool:
		return _es_inicio_pregunta(linea) and not patron_respuesta.match(linea) and not _es_basura(linea)

	def _parsear_item_i(lineas_item: list[str]) -> list[dict[str, str]]:
		registros: list[dict[str, str]] = []
		pregunta_actual: list[str] = []
		opciones_actuales: list[str] = []
		numero = 1

		def _cerrar() -> None:
			nonlocal pregunta_actual, opciones_actuales, numero
			texto = _limpiar_bloque_enunciado(pregunta_actual + opciones_actuales)
			if texto:
				registros.append(
					{
						"codigo": f"I.{numero}",
						"enunciado": texto,
						"display": f"Ítem I.{numero}: {texto}",
					}
				)
				numero += 1
			pregunta_actual = []
			opciones_actuales = []

		for linea in lineas_item:
			if _es_basura(linea):
				continue
			match_num = patron_numero.match(linea)
			match_opcion = patron_opcion.match(linea)
			if match_num and (_es_item_i(linea) or pregunta_actual):
				if pregunta_actual or opciones_actuales:
					_cerrar()
				pregunta_actual = [match_num.group(2).strip() or linea]
				opciones_actuales = []
				continue
			if match_opcion and pregunta_actual:
				opciones_actuales.append(f"{match_opcion.group(1).lower()}) {match_opcion.group(2).strip()}")
				continue
			if patron_vf.match(linea):
				continue
			if not pregunta_actual:
				if "¿" in linea or patron_texto_interrogativo.match(linea):
					pregunta_actual = [linea]
				continue
			if pregunta_actual and not match_opcion:
				pregunta_actual.append(linea)
		_cerrar()
		return registros

	def _parsear_item_ii(lineas_item: list[str]) -> list[dict[str, str]]:
		registros: list[dict[str, str]] = []
		buffer: list[str] = []
		numero = 1
		for linea in lineas_item:
			if _es_basura(linea) or patron_respuesta.match(linea):
				continue
			if patron_vf.match(linea):
				continue
			if patron_bloque_subrayado.match(linea):
				if buffer:
					texto = _limpiar_bloque_enunciado(buffer)
					if texto:
						registros.append({"codigo": f"II.{numero}", "enunciado": texto, "display": f"Ítem II.{numero}: {texto}"})
						numero += 1
					buffer = []
				continue
			texto_linea = linea
			texto_linea = re.sub(r"^_+\s*", "", texto_linea)
			texto_linea = texto_linea.lstrip(". ")
			if texto_linea:
				buffer.append(texto_linea)
				if len(buffer) >= 1:
					texto = _limpiar_bloque_enunciado(buffer)
					if texto:
						registros.append({"codigo": f"II.{numero}", "enunciado": texto, "display": f"Ítem II.{numero}: {texto}"})
						numero += 1
						buffer = []
						continue
		if buffer:
			texto = _limpiar_bloque_enunciado(buffer)
			if texto:
				registros.append({"codigo": f"II.{numero}", "enunciado": texto, "display": f"Ítem II.{numero}: {texto}"})
		return registros

	def _parsear_item_iii(lineas_item: list[str]) -> list[dict[str, str]]:
		registros: list[dict[str, str]] = []
		contexto: list[str] = []
		requerimientos: list[str] = []
		numero = 1
		for linea in lineas_item:
			if _es_basura(linea) or patron_respuesta.match(linea):
				continue
			match_num = patron_numero.match(linea)
			if match_num:
				requerimiento = match_num.group(2).strip() or linea
				requerimiento = _limpiar_enunciado_pregunta(requerimiento)
				if requerimiento:
					requerimientos.append(requerimiento)
				continue
			if not requerimientos:
				contexto.append(linea)
			else:
				requerimientos[-1] = f"{requerimientos[-1]} {linea}".strip()

		contexto_limpio = _limpiar_bloque_enunciado(contexto)
		if not requerimientos and contexto_limpio:
			requerimientos = [contexto_limpio]

		for requerimiento in requerimientos:
			enunciado = requerimiento
			if contexto_limpio and requerimiento != contexto_limpio:
				enunciado = f"{contexto_limpio} {requerimiento}".strip()
			registros.append({"codigo": f"III.{numero}", "enunciado": enunciado, "display": f"Ítem III.{numero}: {enunciado}"})
			numero += 1
		return registros

	registros: list[dict[str, str]] = []
	registros.extend(_parsear_item_i(secciones["I"]))
	registros.extend(_parsear_item_ii(secciones["II"]))
	registros.extend(_parsear_item_iii(secciones["III"]))

	if not registros:
		parrafos = [
			_limpiar_enunciado_pregunta(linea)
			for linea in lineas
			if linea and not _es_linea_portada_instrumento(linea) and not _es_basura(linea)
		]
		return _fallback_por_parrafos(parrafos)

	registros_unicos: list[dict[str, str]] = []
	vistos_codigos: set[str] = set()
	for registro in registros:
		codigo = str(registro.get("codigo", "")).strip().upper()
		if not codigo or codigo in vistos_codigos:
			continue
		vistos_codigos.add(codigo)
		registros_unicos.append(registro)

	return registros_unicos


def _detectar_preguntas_instrumento(texto: str, permitir_secuencial: bool = False) -> list[str]:
	return [item["codigo"] for item in _detectar_preguntas_instrumento_detalladas(texto, permitir_secuencial)]


def _resolver_indicador_asociado(req: str, indicador_texto: str, indicadores: list[str]) -> str:
	if not indicadores:
		return _normalizar_codigo_req(req) or "R4.1.1"

	req_norm = _normalizar_codigo_req(req)  # Ej: "R4.1"
	# Buscar si en el texto del indicador dice "IE2", "IE3", etc.
	match_ie = re.search(r'IE\s*(\d+)', str(indicador_texto), re.IGNORECASE)

	# Si detecta el número de IE, construye el código exacto (Ej: R4.1 + . + 2 = R4.1.2)
	if req_norm and match_ie:
		codigo_construido = f"{req_norm}.{match_ie.group(1)}"
		indicadores_norm = [str(i).strip().upper() for i in indicadores]
		if codigo_construido in indicadores_norm:
			return codigo_construido

	ind_norm = _normalizar_codigo_req(indicador_texto)
	indicadores_norm = [str(i).strip().upper() for i in indicadores]

	# Búsqueda tradicional por coincidencia exacta
	for codigo in indicadores_norm:
		if ind_norm and (codigo == ind_norm or codigo.startswith(ind_norm) or ind_norm.startswith(codigo)):
			return codigo

	# Búsqueda de respaldo
	for codigo in indicadores_norm:
		if req_norm and (codigo == req_norm or codigo.startswith(req_norm)):
			return codigo

	return indicadores_norm[0]


def _normalizar_lista_indicadores(indicadores_texto: str) -> str:
	tokens = re.split(r"[,;\n]+", str(indicadores_texto or ""))
	limpios: list[str] = []
	vistos: set[str] = set()
	for token in tokens:
		codigo = _normalizar_codigo_req(token)
		if not codigo:
			continue
		if codigo in vistos:
			continue
		vistos.add(codigo)
		limpios.append(codigo)
	return ", ".join(limpios)


def _parsear_indicadores_asociados(indicadores_texto: str) -> list[str]:
	tokens = re.split(r"[,;\n]+", str(indicadores_texto or ""))
	resultados: list[str] = []
	vistos: set[str] = set()
	for token in tokens:
		codigo = _normalizar_codigo_req(token)
		if not codigo:
			continue
		if codigo in vistos:
			continue
		vistos.add(codigo)
		resultados.append(codigo)
	return resultados


def _consolidar_pregunta_indicadores(filas: list[dict[str, Any]]) -> list[dict[str, Any]]:
	acumulado: dict[str, dict[str, Any]] = {}
	for fila in filas:
		numero = str(fila.get("numero_pregunta", "")).strip().upper()
		if not numero:
			continue
		try:
			puntos = float(fila.get("puntos_maximos", 0.0))
		except (TypeError, ValueError):
			puntos = 0.0
		indicadores = _parsear_indicadores_asociados(str(fila.get("indicadores_asociados", "")))
		if not indicadores:
			continue

		if numero not in acumulado:
			acumulado[numero] = {"puntos_maximos": puntos, "indicadores": []}
		elif puntos > 0:
			# Conservamos el mayor puntaje si aparece repetido para la misma pregunta.
			acumulado[numero]["puntos_maximos"] = max(float(acumulado[numero]["puntos_maximos"]), puntos)

		for ind in indicadores:
			if ind not in acumulado[numero]["indicadores"]:
				acumulado[numero]["indicadores"].append(ind)

	salida: list[dict[str, Any]] = []
	for numero in sorted(acumulado.keys(), key=_clave_orden_pregunta):
		info = acumulado[numero]
		salida.append(
			{
				"numero_pregunta": numero,
				"puntos_maximos": float(info["puntos_maximos"]),
				"indicadores_asociados": ", ".join(info["indicadores"]),
			}
		)
	return salida


def _extraer_mapeo_docx_tabla_tributacion(
	archivo_bytes: bytes,
	indicadores: list[str],
) -> list[dict[str, Any]]:
	try:
		doc = Document(io.BytesIO(archivo_bytes))
	except Exception:
		return []

	# -- NUEVO: Extracción inteligente de puntajes por ítem --
	mapa_puntos = {}
	texto_completo = "\n".join([p.text for p in doc.paragraphs])
	# Busca patrones como "ITEM I: ... (03 puntos)" tolerando saltos de línea intermedios
	patrones_puntos = re.finditer(r"ITEM\s+([IVX]+)[^\(]{0,100}\(\s*0?(\d+(?:[.,]\d+)?)\s*puntos?", texto_completo, re.IGNORECASE)
	for match in patrones_puntos:
		item_roman = match.group(1).upper()
		try:
			mapa_puntos[item_roman] = float(match.group(2).replace(',', '.'))
		except ValueError:
			pass

	for table in doc.tables:
		filas: list[list[str]] = []
		for row in table.rows:
			fila_textos = [cell.text.strip().replace("\n", " ") for cell in row.cells]
			fila_textos = [re.sub(r"\s+", " ", texto).strip() for texto in fila_textos]
			if any(fila_textos):
				filas.append(fila_textos)

		header_idx = -1
		for i, fila in enumerate(filas):
			norm_fila = [_normalizar_texto_parser(celda) for celda in fila]
			tiene_req = any("req" in celda for celda in norm_fila)
			tiene_pregunta = any("pregunta" in celda for celda in norm_fila)
			if tiene_req and tiene_pregunta:
				header_idx = i
				break

		if header_idx < 0:
			continue

		propuesta_raw: list[dict[str, Any]] = []
		ultimo_req_valido = ""
		for fila in filas[header_idx + 1:]:
			req_celda = fila[0] if len(fila) > 0 else ""
			req = str(req_celda or "").strip()
			if _es_req_valido(req):
				ultimo_req_valido = req
			elif ultimo_req_valido:
				req = ultimo_req_valido
			else:
				continue

			preguntas_txt = fila[-1] if fila else ""
			preguntas = _extraer_preguntas_limpias(preguntas_txt)
			if not preguntas:
				continue

			indicador_txt = fila[4] if len(fila) > 4 else ""
			indicador_asociado = _resolver_indicador_asociado(req, indicador_txt, indicadores)

			for pregunta in preguntas:
				# -- NUEVO: Asignar el puntaje detectado según la letra romana de la pregunta --
				puntos_pregunta = 1.0
				match_item = re.match(r"^([IVX]+)\.", pregunta)
				if match_item and match_item.group(1).upper() in mapa_puntos:
					puntos_pregunta = mapa_puntos[match_item.group(1).upper()]

				propuesta_raw.append(
					{
						"numero_pregunta": pregunta,
						"puntos_maximos": puntos_pregunta,
						"indicadores_asociados": indicador_asociado,
					}
				)

		if propuesta_raw:
			return _consolidar_pregunta_indicadores(propuesta_raw)

	return []


def _extraer_mapeo_pdf_tabla_tributacion(
	archivo_bytes: bytes,
	indicadores: list[str],
) -> list[dict[str, Any]]:
	texto = _leer_texto_pdf_desde_bytes(archivo_bytes)

	# -- NUEVO: Extracción inteligente de puntajes por ítem --
	mapa_puntos = {}
	patrones_puntos = re.finditer(r"ITEM\s+([IVX]+)[^\(]{0,100}\(\s*0?(\d+(?:[.,]\d+)?)\s*puntos?", texto, re.IGNORECASE)
	for match in patrones_puntos:
		item_roman = match.group(1).upper()
		try:
			mapa_puntos[item_roman] = float(match.group(2).replace(',', '.'))
		except ValueError:
			pass

	norm = _normalizar_texto_parser(texto)
	requeridos = ["req", "criterios de evaluacion", "indicador de evaluacion", "pregunta"]
	if not all(token in norm for token in requeridos):
		return []

	propuesta_raw: list[dict[str, Any]] = []
	for linea in texto.splitlines():
		linea_limpia = linea.strip()
		if not linea_limpia:
			continue

		match_req = re.match(r"^\s*(R\s*\d+(?:[\.,]\d+)*)", linea_limpia, flags=re.IGNORECASE)
		if not match_req:
			continue

		req = match_req.group(1)
		if not _es_req_valido(req):
			continue

		preguntas = _extraer_preguntas_limpias(linea_limpia)
		if not preguntas:
			continue

		indicador_asociado = _resolver_indicador_asociado(req, linea_limpia, indicadores)
		for pregunta in preguntas:
			# -- NUEVO: Asignar el puntaje detectado --
			puntos_pregunta = 1.0
			match_item = re.match(r"^([IVX]+)\.", pregunta)
			if match_item and match_item.group(1).upper() in mapa_puntos:
				puntos_pregunta = mapa_puntos[match_item.group(1).upper()]

			propuesta_raw.append(
				{
					"numero_pregunta": pregunta,
					"puntos_maximos": puntos_pregunta,
					"indicadores_asociados": indicador_asociado,
				}
			)

	return _consolidar_pregunta_indicadores(propuesta_raw)


def _normalizar_columna_excel(nombre_columna: Any) -> str:
	texto = str(nombre_columna).strip().lower()
	texto = "".join(
		c for c in unicodedata.normalize("NFKD", texto) if not unicodedata.combining(c)
	)
	texto = re.sub(r"[^a-z0-9]+", " ", texto).strip()
	return re.sub(r"\s+", " ", texto)


def _mapear_columnas_indicadores_excel(df: pd.DataFrame) -> pd.DataFrame | None:
	alias_codigo = {
		"codigo",
		"codigo completo",
		"codigo indicador",
		"codigo del indicador",
		"codigo completo indicador",
	}
	alias_descripcion = {
		"descripcion",
		"descripcion indicador",
		"descripcion del indicador",
		"detalle",
		"detalle indicador",
	}

	candidatas_codigo: list[str] = []
	candidatas_descripcion: list[str] = []

	for col in df.columns:
		norm = _normalizar_columna_excel(col)
		es_codigo = norm in alias_codigo or (
			"codigo" in norm and ("completo" in norm or "indicador" in norm or norm == "codigo")
		)
		es_descripcion = norm in alias_descripcion or "descripcion" in norm or "detalle" in norm

		if es_codigo:
			candidatas_codigo.append(str(col))
		if es_descripcion:
			candidatas_descripcion.append(str(col))

	if len(candidatas_codigo) != 1 or len(candidatas_descripcion) != 1:
		return None

	col_codigo = candidatas_codigo[0]
	col_descripcion = candidatas_descripcion[0]
	if col_codigo == col_descripcion:
		return None

	return df.rename(columns={col_codigo: "Codigo_Completo", col_descripcion: "Descripcion"})


def _unidad_activa() -> dict[str, Any] | None:
	uid = st.session_state.unidad_activa_id
	if not uid:
		return None
	try:
		return obtener_unidad_por_id(int(uid))
	except (TypeError, ValueError):
		return None


def _render_login() -> None:
	st.markdown(
		"""
		<div class="hero">
			<h2>🛡️ Acceso Institucional ACAPOMIL</h2>
			<p>Autenticacion requerida. El sistema se habilita solo con sesion valida.</p>
		</div>
		""",
		unsafe_allow_html=True,
	)

	with st.container(border=True):
		tab_login, tab_registro = st.tabs(["🔐 Iniciar Sesion", "📝 Registrarse"])

		with tab_login:
			st.subheader("Acceso docente")
			with st.form("login_form_docente"):
				usuario = st.text_input("Usuario", key="login_usuario", placeholder="Ingrese su nombre de usuario o correo")
				clave = st.text_input("Clave institucional", type="password", key="login_clave")
				submit_login = st.form_submit_button("Ingresar al sistema", use_container_width=True, type="primary")

				if submit_login:
					# Xuanwu: autenticamos contra SQLite para persistencia real.
					credenciales = obtener_usuario_por_credenciales(usuario.strip(), clave)
					if credenciales:
						st.session_state.is_logged = True
						st.session_state.usuario_activo = credenciales["nombre_completo"]
						st.session_state.username_activo = credenciales["username"]
						st.session_state.docente_id = int(credenciales["id"])
						st.session_state.usuario_rol = str(credenciales.get("rol", "docente"))
						st.session_state.unidad_activa_id = None
						st.rerun()
					else:
						st.error("Credenciales invalidas para acceso docente.")

		with tab_registro:
			st.subheader("Registro de nuevo docente")
			with st.form("registro_docente_form", clear_on_submit=True):
				nuevo_username = st.text_input("Username", placeholder="usuario.docente")
				nuevo_nombre_completo = st.text_input("Nombre completo", placeholder="Nombre Apellido")
				nueva_clave = st.text_input("Clave", type="password")
				registro_ok = st.form_submit_button("Registrar docente", use_container_width=True)

			if registro_ok:
				username = nuevo_username.strip()
				nombre_completo = nuevo_nombre_completo.strip()
				if not username or not nombre_completo or not nueva_clave.strip():
					st.error("Completa username, nombre completo y clave para registrarte.")
				elif obtener_usuario_por_username(username):
					st.warning("Ese usuario ya existe. Usa otro nombre o inicia sesion.")
				else:
					# Xuanwu: registro docente persistente en SQLite.
					crear_usuario(username, nombre_completo, nueva_clave, "docente")
					st.success("Registro exitoso. Ya puedes iniciar sesion en la pestana de acceso.")

	with st.container(border=True):
		st.markdown("### 🎯 Perfil habilitado")
		st.markdown("- Docente por defecto: **123** (123)")
		st.markdown("- Admin por defecto: **admin**")
		st.markdown("- Clave inicial docente por defecto: **123**")
		st.markdown("- Clave inicial admin por defecto: **admin123**")
		st.markdown("- Ambito: Evaluacion por competencias")
		st.markdown("- Seguridad: Sesion obligatoria antes de operar")


def _render_top_context() -> None:
	if st.session_state.usuario_rol == "admin":
		st.markdown(
			"""
			<div class="hero">
				<h2>🛡️ Panel de Administracion ACAPOMIL</h2>
				<p>Control centralizado de matricula semestral por CRIM y Mencion.</p>
			</div>
			""",
			unsafe_allow_html=True,
		)
		st.markdown("**Panel de Control:** admin (Administrador)")
		return

	unidad = _unidad_activa()
	st.markdown(
		"""
		<div class="hero">
			<h2>📚 Contextos de Asignatura</h2>
			<p>Operacion tactica de retroalimentacion RAG por unidad de aprendizaje.</p>
		</div>
		""",
		unsafe_allow_html=True,
	)
	st.markdown(f"**Docente:** {st.session_state.usuario_activo}")

	if unidad:
		st.markdown(
			f"<span class='tag'>Unidad activa: {unidad['nombre']}</span>"
			f"<span class='tag'>Año: {unidad['anio']}</span>"
			f"<span class='tag'>Mencion: {unidad['mencion']}</span>",
			unsafe_allow_html=True,
		)
	else:
		st.warning("No hay Unidad Activa seleccionada. Define una para habilitar el resto del flujo.")


def _menu_sidebar() -> str:
	st.sidebar.title("🧭 Modulos")
	if st.session_state.usuario_rol == "admin":
		st.sidebar.success("Perfil: Administrador")
	else:
		unidad = _unidad_activa()
		if unidad:
			st.sidebar.success(f"Unidad activa: {unidad['nombre']}")
		else:
			st.sidebar.info("Sin unidad activa")

	if st.sidebar.button("Cerrar sesion", use_container_width=True):
		st.session_state.is_logged = False
		st.session_state.usuario_activo = ""
		st.session_state.username_activo = ""
		st.session_state.docente_id = None
		st.session_state.usuario_rol = ""
		st.session_state.unidad_activa_id = None
		st.rerun()

	if st.session_state.usuario_rol == "admin":
		return st.sidebar.radio(
			"Navegacion",
			[
				"1) Control de Matricula",
				"2) Gestion de Cuentas Docentes",
			],
		)

	return st.sidebar.radio(
		"Navegacion",
		[
			"1) Gestion de Unidades",
			"2) Bibliografia y Vectorizacion",
			"3) Diseño de Evaluación",
			"4) Trazabilidad RAG",
			"5) Reportes Jefatura",
		],
	)


def _vista_admin_matricula() -> None:
	st.subheader("🗂️ Control de Matricula Semestral")
	with st.container(border=True):
		st.info("Instrucciones de uso: administra altas y bajas de alumnos por CRIM y Mencion. Este panel es exclusivo para Administrador.")

	with st.container(border=True):
		st.markdown("### ➕ Matricular alumno")
		c1, c2 = st.columns(2)
		with c1:
			nombre = st.text_input("Nombre completo", key="admin_nombre")
			rut = st.text_input("RUT", placeholder="12.345.678-9", key="admin_rut")
			email = st.text_input("Correo Institucional", placeholder="alumno@acapomil.cl", key="admin_email")
		with c2:
			anio = st.selectbox("Año", ANIOS_CRIM, key="admin_anio")
			mencion = st.selectbox("Mencion", MENCIONES, key="admin_mencion")

		if st.button("Guardar matricula", use_container_width=True, type="primary"):
			rut_ingresado = rut.strip()
			if not nombre.strip() or not rut_ingresado:
				st.error("Nombre y RUT son obligatorios.")
			elif not validar_rut_chileno(rut_ingresado):
				st.error("❌ El RUT ingresado es inválido según el algoritmo estándar chileno. Por favor, verifíquelo.")
			else:
				rut_formateado = formatear_rut_chileno(rut_ingresado)
				if buscar_alumno_por_rut_anio_mencion(rut=rut_formateado, anio=anio, mencion=mencion):
					st.warning("Ese alumno ya esta matriculado para ese CRIM y Mencion.")
				else:
					# Xuanwu: matricula centralizada persistente para consumo automatico docente.
					crear_alumno(rut=rut_formateado, nombre=nombre.strip(), email=email.strip(), anio=anio, mencion=mencion)
					st.success("Alumno matriculado correctamente.")
					st.rerun()

	with st.container(border=True):
		st.markdown("### 🧾 Matricula registrada")
		filtro_anio = st.selectbox("Filtrar por Año", ["Todos"] + ANIOS_CRIM, key="admin_filtro_anio")
		filtro_mencion = st.selectbox("Filtrar por Mencion", ["Todas"] + MENCIONES, key="admin_filtro_mencion")
		anio_sel = None if filtro_anio == "Todos" else filtro_anio
		mencion_sel = None if filtro_mencion == "Todas" else filtro_mencion
		alumnos = listar_alumnos(anio=anio_sel, mencion=mencion_sel)

		if not alumnos:
			st.info("No hay alumnos matriculados para el filtro seleccionado.")
			return

		tabla = pd.DataFrame(alumnos)
		tabla = tabla[["rut", "nombre", "email", "anio", "mencion"]]
		st.dataframe(tabla, use_container_width=True)

		mapa = {f"{a['nombre']} ({a['rut']}) | {a['anio']} | {a['mencion']}": int(a["id"]) for a in alumnos}
		obj = st.selectbox("Alumno a eliminar", list(mapa.keys()))
		if st.button("Eliminar matricula", use_container_width=True):
			eliminar_alumno(mapa[obj])
			st.success("Matricula eliminada.")
			st.rerun()


def _vista_admin_cuentas() -> None:
	st.subheader("👤 Gestion de Cuentas Docentes")
	with st.container(border=True):
		st.info("Instrucciones de uso: crea, actualiza o elimina cuentas. Por seguridad, la cuenta admin no puede eliminarse.")

	usuarios = obtener_usuarios()
	with st.container(border=True):
		st.markdown("### 📋 Cuentas registradas")
		if not usuarios:
			st.info("No hay usuarios registrados.")
		else:
			tabla = pd.DataFrame(usuarios)
			tabla = tabla[["nombre_completo", "username", "rol"]]
			st.dataframe(tabla, use_container_width=True)

	with st.container(border=True):
		st.markdown("### ➕ Crear cuenta docente")
		col_crear_1, col_crear_2 = st.columns(2)
		with col_crear_1:
			new_username = st.text_input("Username")
			new_nombre = st.text_input("Nombre Completo (con Grado Militar)")
		with col_crear_2:
			new_clave = st.text_input("Clave", type="password")
			new_rol = st.selectbox("Asignar Rol", ["docente", "admin"])

		if st.button("Registrar cuenta", use_container_width=True):
			if not new_username or not new_clave or not new_nombre:
				st.warning("⚠️ Por favor, rellene todos los campos antes de registrar.")
			else:
				try:
					# Intentamos la insercion en la base de datos
					crear_usuario(
						username=new_username.strip(),
						nombre_completo=new_nombre.strip(),
						clave=new_clave.strip(),
						rol=new_rol,
					)
					st.success("🎉 ¡Cuenta registrada con éxito!")
					st.rerun()
				except Exception as e:
					st.error(f"❌ Error crítico en el backend: {str(e)}")

	with st.container(border=True):
		st.markdown("### ✏️ Editar o eliminar cuenta")
		if not usuarios:
			st.info("Sin cuentas para administrar.")
			return

		usernames = [u["username"] for u in usuarios]
		username_objetivo = st.selectbox("Cuenta objetivo", usernames, key="admin_user_target")
		objetivo = next(u for u in usuarios if u["username"] == username_objetivo)

		col1, col2 = st.columns(2)
		with col1:
			ed_username = st.text_input("Nuevo username", value=objetivo["username"], key="admin_user_edit_username")
			ed_nombre_completo = st.text_input("Nuevo nombre completo", value=objetivo["nombre_completo"], key="admin_user_edit_fullname")
		with col2:
			ed_clave = st.text_input("Nueva clave", type="password", key="admin_user_edit_pwd", placeholder="Dejar vacio para mantener")
			opciones_rol = ["docente", "admin"]
			rol_actual = str(objetivo.get("rol", "docente")).strip().lower()
			rol_index = opciones_rol.index(rol_actual) if rol_actual in opciones_rol else 0
			ed_rol = st.selectbox(
				"Rol",
				opciones_rol,
				index=rol_index,
				key=f"admin_user_edit_role_{username_objetivo}",
			)

		confirmar = st.checkbox("Confirmar eliminacion", key="admin_user_delete_confirm")

		act_col1, act_col2 = st.columns(2)
		with act_col1:
			if st.button("Guardar cambios cuenta", use_container_width=True):
				username_final = ed_username.strip() or objetivo["username"]
				nombre_completo_final = ed_nombre_completo.strip() or objetivo["nombre_completo"]
				existente = obtener_usuario_por_username(username_final)
				if username_final != objetivo["username"] and existente:
					st.warning("Ya existe otra cuenta con ese nombre.")
				elif objetivo["username"] == "admin" and ed_rol != "admin":
					st.error("La cuenta admin no puede cambiar su rol a docente.")
				else:
					try:
						actualizar_usuario(
							usuario_id=int(objetivo["id"]),
							username=username_final,
							nombre_completo=nombre_completo_final,
							clave=ed_clave if ed_clave.strip() else None,
							rol=ed_rol,
						)
						st.success("Cuenta actualizada.")
						st.rerun()
					except ValueError as exc:
						st.error(str(exc))

		with act_col2:
			if st.button("Eliminar cuenta", use_container_width=True, type="secondary"):
				if not confirmar:
					st.warning("Debes confirmar antes de eliminar.")
				elif username_objetivo == "admin" or int(objetivo["id"]) == int(st.session_state.docente_id or -1):
					st.error("No puedes eliminar la cuenta administradora activa.")
				else:
					try:
						# Xuanwu: eliminacion robusta por username unico seleccionado en el selectbox.
						eliminar_usuario_por_username(username_objetivo)
						st.success("Cuenta eliminada.")
						st.rerun()
					except ValueError as exc:
						st.error(str(exc))
					except sqlite3.IntegrityError:
						st.error("No se puede eliminar: Este docente tiene unidades de aprendizaje asignadas en el sistema.")
					except Exception:
						st.error("Error al intentar eliminar esta cuenta. Verifica las dependencias en la base de datos.")


def _vista_unidades() -> None:
	st.subheader("🏛️ Gestion de Unidades de Aprendizaje")
	with st.container(border=True):
		st.info("Instrucciones de uso: crea una unidad, activala y luego edita o elimina segun necesidad academica.")

	docente_id = st.session_state.docente_id
	if docente_id is None:
		st.error("Sesion docente no valida. Vuelve a iniciar sesion.")
		return

	unidades = listar_unidades(docente_id=docente_id)

	with st.expander("➕ Crear nueva unidad de aprendizaje", expanded=False):
		with st.form("crear_unidad_form", clear_on_submit=True):
			nombre = st.text_input("Nombre de unidad", placeholder="Bases de Datos")
			anio = st.selectbox("Año", ANIOS_CRIM)
			mencion = st.selectbox("Mencion", MENCIONES)
			submitted = st.form_submit_button("Crear unidad", use_container_width=True)

		if submitted:
			if not nombre.strip():
				st.error("Ingresa un nombre valido para la unidad.")
			else:
				# Xuanwu: alta de unidad persistente en base relacional.
				unidad_id = crear_unidad(nombre.strip(), anio, mencion, docente_id)
				if not st.session_state.unidad_activa_id:
					st.session_state.unidad_activa_id = unidad_id
				# Xuanwu: preparamos carpetas por unidad para doctrina y vector store.
				_data_unidad_dir(str(unidad_id)).mkdir(parents=True, exist_ok=True)
				_vector_unidad_dir(str(unidad_id)).mkdir(parents=True, exist_ok=True)
				st.success("Unidad creada y disponible para seleccion.")
				st.rerun()

	with st.container(border=True):
		st.markdown("### 🎛️ Seleccionar unidad activa")
		if not unidades:
			st.info("Aun no existen unidades.")
		else:
			mapa = {
				f"{u['nombre']} | {u['anio']} | {u['mencion']}": u["id"]
				for u in unidades
			}
			opciones = list(mapa.keys())
			idx = 0
			if st.session_state.unidad_activa_id is not None:
				for i, key in enumerate(opciones):
					if mapa[key] == st.session_state.unidad_activa_id:
						idx = i
						break
			seleccion = st.selectbox("Unidad", opciones, index=idx)
			if st.button("Activar unidad", use_container_width=True):
				st.session_state.unidad_activa_id = mapa[seleccion]
				st.success("Unidad activa actualizada.")
				# Xuanwu: refresco explicito para actualizar cabecera y sidebar de inmediato.
				st.rerun()

	with st.expander("🛠️ Editar / Eliminar unidad", expanded=False):
		if not unidades:
			st.info("Sin unidades para administrar.")
			return

		mapa = {f"{u['nombre']} ({u['anio']})": u for u in unidades}
		pick = st.selectbox("Unidad objetivo", list(mapa.keys()), key="editar_unidad_pick")
		target = mapa[pick]

		ncol1, ncol2, ncol3 = st.columns(3)
		with ncol1:
			nuevo_nombre = st.text_input("Nuevo nombre", value=target["nombre"], key="ed_nombre")
		with ncol2:
			nuevo_anio = st.selectbox(
				"Nuevo año",
				ANIOS_CRIM,
				index=ANIOS_CRIM.index(target["anio"]),
				key="ed_anio",
			)
		with ncol3:
			nueva_mencion = st.selectbox(
				"Nueva mencion",
				MENCIONES,
				index=MENCIONES.index(target["mencion"]),
				key="ed_mencion",
			)

		a, b = st.columns(2)
		with a:
			if st.button("Guardar cambios", use_container_width=True):
				actualizar_unidad(
					unidad_id=int(target["id"]),
					nombre=nuevo_nombre.strip() or target["nombre"],
					anio=nuevo_anio,
					mencion=nueva_mencion,
				)
				st.success("Unidad actualizada.")
				st.rerun()
		with b:
			confirmar = st.checkbox("Confirmar eliminacion", key="confirm_del_unit")
			if st.button("Eliminar unidad", use_container_width=True, type="secondary"):
				if not confirmar:
					st.warning("Marca la confirmacion para eliminar.")
				else:
					uid = int(target["id"])
					eliminar_unidad(uid)
					shutil.rmtree(_data_unidad_dir(str(uid)), ignore_errors=True)
					shutil.rmtree(_vector_unidad_dir(str(uid)), ignore_errors=True)
					if st.session_state.unidad_activa_id == uid:
						restantes = listar_unidades(docente_id=docente_id)
						st.session_state.unidad_activa_id = restantes[0]["id"] if restantes else None
					st.success("Unidad eliminada.")
					st.rerun()


def _vista_diseno_evaluacion() -> None:
	st.subheader("🧩 Diseño y Mapeo de Evaluaciones")
	with st.container(border=True):
		st.info("Instrucciones de uso: completa los 3 pasos del asistente para definir indicadores, construir la matriz base y validar la cobertura doctrinal.")

	unidad = _unidad_activa()
	if not unidad:
		st.info("Activa una unidad para habilitar el diseño de evaluacion.")
		return

	unidad_id = int(unidad["id"])
	st.session_state.wizard_eval_step = st.session_state.get("wizard_eval_step", 1)

	pasos = {1: "Paso 1: Indicadores", 2: "Paso 2: Matriz Base", 3: "Paso 3: Asistente IA"}
	nav_col1, nav_col2, nav_col3 = st.columns(3)
	for columna, paso_num in zip((nav_col1, nav_col2, nav_col3), (1, 2, 3)):
		with columna:
			if st.session_state.wizard_eval_step == paso_num:
				st.markdown(f"**➡️ {pasos[paso_num]}**")
			else:
				st.markdown(pasos[paso_num])

	# ---------------- Paso 1: Indicadores ----------------
	if st.session_state.wizard_eval_step == 1:
		with st.container(border=True):
			st.markdown("### Fase 1: Indicadores de la Unidad Activa")
			cargador_col, info_col = st.columns(2)
			with cargador_col:
				archivo_indicadores = st.file_uploader(
					"Subir Matriz de Indicadores (.xlsx / .pdf)",
					type=["xlsx", "pdf"],
					key=f"excel_uploader_{st.session_state.uploader_key}",
				)
			with info_col:
				st.caption("Excel: columnas Codigo_Completo y Descripcion. PDF: texto con codigo y descripcion.")

			if archivo_indicadores is not None:
				evento_actual = f"{unidad_id}:{archivo_indicadores.name}:{archivo_indicadores.size}:{st.session_state.uploader_key}"
				if st.session_state.indicadores_excel_evento != evento_actual:
					st.session_state.indicadores_excel_evento = evento_actual
					st.session_state.indicadores_excel_procesado = False

				if not st.session_state.indicadores_excel_procesado:
					st.session_state.indicadores_excel_procesado = True
					if archivo_indicadores.name.lower().endswith(".pdf"):
						try:
							indicadores_pdf = extraer_indicadores_desde_pdf(archivo_indicadores.getvalue())
						except Exception as exc:
							st.error(f"❌ Error al leer el PDF: {exc}")
						else:
							indicadores_limpios: list[dict[str, str]] = []
							vistos: set[str] = set()
							for fila in indicadores_pdf:
								codigo = str(fila.get("codigo", "")).strip().upper()
								descripcion = str(fila.get("descripcion", "")).strip()
								if not codigo or codigo in vistos:
									continue
								vistos.add(codigo)
								indicadores_limpios.append({"codigo": codigo, "descripcion": descripcion})

							if not indicadores_limpios:
								st.error("No se encontraron indicadores válidos en el PDF.")
							else:
								reemplazar_indicadores_unidad(unidad_id, indicadores_limpios)
								st.success(f"🎉 ¡Se extrajeron y guardaron {len(indicadores_limpios)} indicadores desde el PDF!")
								st.session_state.uploader_key += 1
								st.session_state.indicadores_excel_evento = ""
								st.session_state.indicadores_excel_procesado = False
								st.rerun()
					else:
						try:
							df_indicadores = pd.read_excel(archivo_indicadores)
						except Exception as exc:
							st.error(f"❌ Error al leer el archivo Excel: {exc}")
						else:
							df_mapeado = _mapear_columnas_indicadores_excel(df_indicadores)
							if df_mapeado is None:
								st.error("❌ Formato de Excel inválido. Debe contener las columnas 'Codigo_Completo' y 'Descripcion'.")
							else:
								cargar_indicadores_desde_df(unidad_id, df_mapeado)
								st.success(f"🎉 ¡Se cargaron con éxito {len(df_mapeado)} indicadores para esta Unidad de Aprendizaje!")
								st.session_state.uploader_key += 1
								st.session_state.indicadores_excel_evento = ""
								st.session_state.indicadores_excel_procesado = False
								st.rerun()
			else:
				indicadores_db = listar_indicadores_unidad(unidad_id)

				# Mostrar feedback claro al usuario sobre el estado actual
				if indicadores_db:
					st.success(f"✅ Hay {len(indicadores_db)} indicadores guardados en la base de datos para esta Unidad.")
				else:
					st.warning("⚠️ No hay indicadores guardados. Sube un archivo o créalos manualmente en la tabla inferior.")

				base_rows = [
					{"codigo": row["codigo"], "descripcion": row.get("descripcion", "") or ""}
					for row in indicadores_db
				] or [{"codigo": "R4.1.1", "descripcion": ""}]

				editor_indicadores = st.data_editor(
					pd.DataFrame(base_rows),
					use_container_width=True,
					num_rows="dynamic",
					key=f"editor_indicadores_unidad_{unidad_id}",
				)

				if st.button("Guardar indicadores de unidad", use_container_width=True):
					filas = editor_indicadores.to_dict(orient="records")
					codigos = [str(f.get("codigo", "")).strip() for f in filas if str(f.get("codigo", "")).strip()]
					if not codigos:
						st.error("Debes definir al menos un indicador para la unidad.")
					elif len(codigos) != len(set(codigos)):
						st.error("Existen codigos de indicadores duplicados. Corrigelos antes de guardar.")
					else:
						reemplazar_indicadores_unidad(unidad_id, filas)
						st.success("Indicadores de unidad actualizados.")
						st.rerun()

			if st.button("Siguiente: Crear Matriz Base", use_container_width=True, type="primary"):
				st.session_state.wizard_eval_step = 2
				st.rerun()

	# ---------------- Paso 2: Matriz Base ----------------
	elif st.session_state.wizard_eval_step == 2:
		st.markdown("### Fase 2: Matriz Base de la Evaluación")
		nombre_eval = st.text_input("Nombre de evaluación", value=st.session_state.get("nombre_eval_temporal", "Solemne 1"))

		# --- INICIO BLOQUE OPCIONAL AUTOCOMPLETADO ---
		with st.expander("📥 (Opcional) Autocompletar matriz desde archivo de evaluación", expanded=False):
			st.info("💡 Si ya tienes tu evaluación en Word o PDF con una tabla de tributación (Req, Indicador, Pregunta), súbela aquí. El sistema extraerá y autocompletará la matriz por ti.")
			archivo_eval = st.file_uploader("Subir evaluación (.docx o .pdf)", type=["docx", "pdf"], key="uploader_autocompletar")

			if archivo_eval and st.button("Procesar y Autocompletar", type="primary"):
				with st.spinner("Analizando documento y extrayendo tabla de tributación..."):
					bytes_archivo = archivo_eval.getvalue()
					indicadores_db = listar_indicadores_unidad(unidad_id)
					codigos_ind = [r["codigo"] for r in indicadores_db]
					opciones_ind_full = [f"{r['codigo']} - {r.get('descripcion', '')}" for r in indicadores_db]

					# Llamada a las funciones extractoras nativas del sistema
					if archivo_eval.name.lower().endswith(".docx"):
						datos_extraidos = _extraer_mapeo_docx_tabla_tributacion(bytes_archivo, codigos_ind)
					else:
						datos_extraidos = _extraer_mapeo_pdf_tabla_tributacion(bytes_archivo, codigos_ind)

					if datos_extraidos:
						nueva_matriz = []
						for d in datos_extraidos:
							# Mapear los códigos extraídos a las opciones completas del multiselect
							cods = [c.strip() for c in str(d.get("indicadores_asociados", "")).split(",")]
							sel_full = [opt for opt in opciones_ind_full if any(opt.startswith(c) for c in cods)]
							nueva_matriz.append({
								"numero_pregunta": d["numero_pregunta"],
								"puntos_maximos": float(d["puntos_maximos"]),
								"indicadores_asociados": sel_full
							})
						st.session_state.matriz_temporal = nueva_matriz
						st.session_state.num_preg_temporal = len(nueva_matriz)
						st.success(f"✅ ¡Éxito! Se autocompletaron {len(nueva_matriz)} preguntas desde el documento.")
						st.rerun()
					else:
						st.warning("⚠️ No se encontró una tabla de tributación válida o el formato no coincide. Puedes armar la matriz manualmente a continuación.")
		# --- FIN BLOQUE OPCIONAL AUTOCOMPLETADO ---

		num_preg = st.number_input("Cantidad de preguntas", min_value=1, max_value=150, value=st.session_state.get("num_preg_temporal", 10))

		indicadores_db = listar_indicadores_unidad(unidad_id)
		opciones_ind = [f"{r['codigo']} - {r.get('descripcion', '')}" for r in indicadores_db]

		st.markdown("#### Asignación de Indicadores por Pregunta")
		col_p, col_pts, col_ind = st.columns([1, 1, 4])
		col_p.markdown("**Pregunta**")
		col_pts.markdown("**Puntos**")
		col_ind.markdown("**Indicadores Asociados (Selección Múltiple)**")

		matriz_guardada = st.session_state.get("matriz_temporal", [])
		matriz_base = []

		for i in range(num_preg):
			# Si hay datos autocompletados, conservamos el nombre original (ej. "I.1", "II.2"). Si es manual, usamos "P1", "P2".
			if i < len(matriz_guardada) and str(matriz_guardada[i].get("numero_pregunta", "")).strip() != "":
				pregunta_id = matriz_guardada[i]["numero_pregunta"]
				fila_previa = matriz_guardada[i]
			else:
				pregunta_id = f"P{i+1}"
				fila_previa = next((f for f in matriz_guardada if f["numero_pregunta"] == pregunta_id), None)

			cp, cpts, cind = st.columns([1, 1, 4])
			cp.markdown(f"**{pregunta_id}**")

			def_pts = float(fila_previa["puntos_maximos"]) if fila_previa else 1.0
			def_ind = [ind for ind in fila_previa["indicadores_asociados"] if ind in opciones_ind] if fila_previa else []

			pts = cpts.number_input("Pts", value=def_pts, step=0.5, key=f"pts_{i}_w2", label_visibility="collapsed")
			sel = cind.multiselect("Ind", options=opciones_ind, default=def_ind, key=f"ind_{i}_w2", label_visibility="collapsed")

			matriz_base.append({
				"numero_pregunta": pregunta_id,
				"puntos_maximos": pts,
				"indicadores_asociados": sel
			})

		c_back, c_next = st.columns(2)
		with c_back:
			if st.button("⬅️ Atrás (Paso 1)"):
				st.session_state.wizard_eval_step = 1
				st.rerun()
		with c_next:
			if st.button("Siguiente: Validar Cobertura ➡️", type="primary"):
				# Validar que al menos una pregunta tenga un indicador asociado
				total_indicadores = sum(len(fila["indicadores_asociados"]) for fila in matriz_base)

				if total_indicadores == 0:
					st.error("⚠️ **Matriz Vacía:** No has asignado ningún indicador a las preguntas. Despliega el menú opcional de arriba para autocompletar con un archivo, o selecciona los indicadores manualmente en la tabla antes de avanzar.")
				else:
					st.session_state.matriz_temporal = matriz_base
					st.session_state.nombre_eval_temporal = nombre_eval
					st.session_state.num_preg_temporal = num_preg
					st.session_state.wizard_eval_step = 3
					st.rerun()

	# ---------------- Paso 3: Validar Cobertura ----------------
	elif st.session_state.wizard_eval_step == 3:
		st.markdown("### Fase 3: Validación de Cobertura Doctrinal")
		matriz = st.session_state.get("matriz_temporal", [])
		indicadores_db = listar_indicadores_unidad(unidad_id)

		indicadores_usados = set()
		for fila in matriz:
			for ind_str in fila["indicadores_asociados"]:
				codigo = ind_str.split(" - ")[0].strip()
				indicadores_usados.add(codigo)

		codigos_totales = set([r["codigo"] for r in indicadores_db])
		faltantes = codigos_totales - indicadores_usados

		if faltantes:
			st.error(f"⚠️ Alerta de Cobertura Doctrinal: Faltan preguntas para evaluar los siguientes indicadores: {', '.join(sorted(faltantes))}")
			for cod in sorted(faltantes):
				if st.button(f"🤖 Generar sugerencia con IA para {cod}"):
					with st.spinner(f"Consultando doctrina y generando pregunta para {cod}..."):
						try:
							sugerencia = generar_sugerencia_pregunta_cobertura(cod, str(unidad_id))
							st.info(sugerencia)
						except Exception as e:
							st.error(f"Error de ejecución IA: {e}")
		else:
			st.success("✅ Matriz 100% Doctrinalmente Coherente. Todos los indicadores están cubiertos.")

		c_back, c_save = st.columns(2)
		with c_back:
			if st.button("⬅️ Atrás (Editar Matriz)"):
				st.session_state.wizard_eval_step = 2
				st.rerun()
		with c_save:
			if st.button("💾 Guardar Evaluación en Base de Datos", type="primary"):
				filas_db = []
				for fila in matriz:
					codigos = [s.split(" - ")[0].strip() for s in fila["indicadores_asociados"]]
					if codigos:
						filas_db.append({
							"numero_pregunta": fila["numero_pregunta"],
							"puntos_maximos": fila["puntos_maximos"],
							"indicadores_asociados": ", ".join(codigos)
						})
				if not filas_db:
					st.error("No puedes guardar una evaluación sin indicadores asignados.")
				else:
					eval_id = guardar_evaluacion(unidad_id, st.session_state.get("nombre_eval_temporal", "Solemne"), "generado_por_app.txt")
					reemplazar_mapeo_evaluacion(eval_id, filas_db)
					st.success("Evaluación y Matriz guardadas con éxito.")
					st.session_state.wizard_eval_step = 1
					st.rerun()



def _reindexar_unidad(unidad_id: str) -> None:
	doc_dir = _data_unidad_dir(unidad_id)
	vector_dir = _vector_unidad_dir(unidad_id)
	doc_dir.mkdir(parents=True, exist_ok=True)

	docs = sorted(
		p
		for p in doc_dir.iterdir()
		if p.is_file() and p.suffix.lower() in {".docx", ".pdf", ".pptx", ".ppt"}
	)
	if not docs:
		shutil.rmtree(vector_dir, ignore_errors=True)
		return

	documentos = cargar_documentos_docx(origen=doc_dir)
	construir_vector_store(
		documentos=documentos,
		persist_directory=vector_dir,
		collection_name=f"acapomil_{unidad_id}",
	)


def _vista_bibliografia() -> None:
	st.subheader("📖 Bibliografia dinamica por Unidad Activa")
	with st.container(border=True):
		st.info("Instrucciones de uso: carga archivos DOCX, PDF o PowerPoint, vectoriza y administra el repositorio documental de la unidad activa.")
	unidad = _unidad_activa()
	if not unidad:
		st.info("Define y activa una unidad para gestionar doctrina.")
		return

	unidad_id = str(unidad["id"])
	doc_dir = _data_unidad_dir(unidad_id)
	doc_dir.mkdir(parents=True, exist_ok=True)

	with st.container(border=True):
		st.markdown("### ⬆️ Cargar documentos doctrinales")
		archivos = st.file_uploader(
			"Cargar doctrina",
			type=["docx", "pdf", "pptx", "ppt"],
			accept_multiple_files=True,
			key=f"upload_{unidad_id}",
		)
		if st.button("Guardar y vectorizar", use_container_width=True, type="primary"):
			if not archivos:
				st.warning("Selecciona al menos un documento DOCX, PDF, PPTX o PPT")
			else:
				for archivo in archivos:
					destino = doc_dir / archivo.name
					destino.write_bytes(archivo.getvalue())
				with st.spinner("Vectorizando bibliografia de la unidad activa..."):
					try:
						_reindexar_unidad(unidad_id)
						st.success("Documentos guardados y vector store actualizado.")
					except Exception as exc:
						st.error(f"No se pudo vectorizar: {exc}")

	docs = sorted(
		p
		for p in doc_dir.iterdir()
		if p.is_file() and p.suffix.lower() in {".docx", ".pdf", ".pptx", ".ppt"}
	)
	with st.container(border=True):
		st.markdown("### 🗂️ Documentos actuales")
		if not docs:
			st.info("No hay documentos cargados para esta unidad.")
		else:
			tabla = pd.DataFrame(
				[
					{
						"Documento": p.name,
						"Tamano (KB)": round(p.stat().st_size / 1024, 2),
					}
					for p in docs
				]
			)
			st.dataframe(tabla, use_container_width=True)

			archivo_del = st.selectbox("Documento a eliminar", [p.name for p in docs])
			c1, c2 = st.columns(2)
			with c1:
				if st.button("Eliminar documento", use_container_width=True):
					(doc_dir / archivo_del).unlink(missing_ok=True)
					try:
						_reindexar_unidad(unidad_id)
					except Exception as exc:
						st.error(f"Documento eliminado, pero fallo reindexacion: {exc}")
					st.success("Documento eliminado.")
					st.rerun()
			with c2:
				if st.button("Eliminar vector store de la unidad", use_container_width=True):
					shutil.rmtree(_vector_unidad_dir(unidad_id), ignore_errors=True)
					st.warning("Vector store eliminado para la unidad activa.")


def _alumnos_unidad(unidad_id: str) -> list[dict[str, Any]]:
	unidad = obtener_unidad_por_id(int(unidad_id)) if str(unidad_id).isdigit() else None
	if not unidad:
		return []
	# Xuanwu: consumo docente centralizado por filtro exacto CRIM + Mencion.
	return listar_alumnos(anio=unidad["anio"], mencion=unidad["mencion"])


def _vista_alumnos() -> None:
	st.subheader("🧑‍🎓 Alumnos de la Unidad Activa")
	with st.container(border=True):
		st.info("Instrucciones de uso: visualiza los alumnos matriculados por Administracion para el CRIM y Mencion de la unidad activa.")
	unidad = _unidad_activa()
	if not unidad:
		st.info("Activa una unidad para visualizar alumnos filtrados automaticamente.")
		return
	alumnos = _alumnos_unidad(str(unidad["id"]))
	with st.container(border=True):
		st.markdown("### 📋 Alumnos de la unidad activa")
		if not alumnos:
			st.info("Aun no hay alumnos registrados en esta unidad.")
			return

		# Xuanwu: mostramos solo columnas pedagogicas, ocultando ids tecnicos.
		tabla_alumnos = pd.DataFrame(alumnos)
		if not tabla_alumnos.empty:
			tabla_alumnos = tabla_alumnos[["rut", "nombre", "anio", "mencion"]]
		st.dataframe(tabla_alumnos, use_container_width=True)
		st.caption("Gestion de altas/bajas restringida al Administrador.")


def _default_result_rows() -> list[dict[str, Any]]:
	return [
		{"Indicador": "R4.1.1", "Logro": 0.0, "Observacion": ""},
		{"Indicador": "R4.2.1", "Logro": 0.0, "Observacion": ""},
	]


def _vista_trazabilidad() -> None:
	st.subheader("🧠 Flujo de Trazabilidad RAG")
	with st.container(border=True):
		st.info("Instrucciones de uso: ingresa puntajes por pregunta de la evaluacion activa, calcula logro por indicador y prioriza brechas para el plan RAG.")
	unidad = _unidad_activa()
	if not unidad:
		st.info("Activa una unidad para habilitar trazabilidad.")
		return

	evaluaciones = listar_evaluaciones_unidad(int(unidad["id"]))
	if not evaluaciones:
		st.warning("No existen evaluaciones diseñadas. Usa primero la pestaña Diseño de Evaluación.")
		return

	mapa_eval = {f"{e['nombre']} (ID {e['id']})": e for e in evaluaciones}
	opciones_eval = list(mapa_eval.keys())
	idx_eval = 0
	for i, key in enumerate(opciones_eval):
		if int(mapa_eval[key]["id"]) == int(st.session_state.evaluacion_activa_id or -1):
			idx_eval = i
			break
	seleccion_eval = st.selectbox("Evaluacion para trazabilidad", opciones_eval, index=idx_eval)
	evaluacion = mapa_eval[seleccion_eval]
	evaluacion_id = int(evaluacion["id"])
	st.session_state.evaluacion_activa_id = evaluacion_id

	mapeo_eval = listar_mapeo_evaluacion(evaluacion_id)
	if not mapeo_eval:
		st.warning("La evaluacion seleccionada no tiene matriz validada. Completa la Fase 3 en Diseño de Evaluación.")
		return
	mapeo_eval = _consolidar_pregunta_indicadores([
		{
			"numero_pregunta": row["numero_pregunta"],
			"puntos_maximos": row["puntos_maximos"],
			"indicadores_asociados": row.get("indicadores_asociados", ""),
		}
		for row in mapeo_eval
	])

	alumnos = listar_alumnos(anio=unidad["anio"], mencion=unidad["mencion"])
	if not alumnos:
		st.warning("Registra alumnos primero para continuar.")
		return

	st.markdown("### 1) Ingreso Masivo de Puntajes (Matriz de Curso)")
	st.info("Ingresa los puntajes obtenidos por cada alumno. La matriz se guardará para toda la sección de forma simultánea.")

	preguntas_ordenadas = sorted(mapeo_eval, key=lambda x: _clave_orden_pregunta(x["numero_pregunta"]))

	# Construir datos para el DataFrame
	datos_curso = []
	for alumno in alumnos:
		alumno_id = int(alumno["id"])
		puntajes_historicos = obtener_puntajes_pregunta(alumno_id=alumno_id, evaluacion_id=evaluacion_id)

		fila = {
			"ID_Alumno": alumno_id,
			"RUT": alumno["rut"],
			"Alumno": alumno["nombre"]
		}
		for p in preguntas_ordenadas:
			num_preg = p["numero_pregunta"]
			col_name = f"{num_preg} (Max: {p['puntos_maximos']})"
			fila[col_name] = float(puntajes_historicos.get(num_preg, 0.0))
		datos_curso.append(fila)

	df_curso = pd.DataFrame(datos_curso)

	# Configurar columnas para st.data_editor
	config_columnas = {
		"ID_Alumno": None, # Ocultar columna técnica
		"RUT": st.column_config.TextColumn(disabled=True, width="medium"),
		"Alumno": st.column_config.TextColumn(disabled=True, width="medium"),
	}
	for p in preguntas_ordenadas:
		col_name = f"{p['numero_pregunta']} (Max: {p['puntos_maximos']})"
		config_columnas[col_name] = st.column_config.NumberColumn(
			min_value=0.0, max_value=float(p['puntos_maximos']), step=0.5, width="small"
		)

	# Sincronización bidireccional con session_state para evitar pérdida de datos por recargas
	clave_editor = f"estado_matriz_{evaluacion_id}"

	def sincronizar_matriz():
		st.session_state[clave_editor] = st.session_state[f"editor_masivo_{evaluacion_id}"]

	if clave_editor not in st.session_state:
		st.session_state[clave_editor] = df_curso
	else:
		# Si la BD cambió estructuralmente (ej. más alumnos), reseteamos el estado. Si no, mantenemos lo que el usuario estaba editando.
		if len(st.session_state[clave_editor]) != len(df_curso):
			st.session_state[clave_editor] = df_curso

	edited_df = st.data_editor(
		st.session_state[clave_editor],
		use_container_width=True,
		hide_index=True,
		column_config=config_columnas,
		key=f"editor_masivo_{evaluacion_id}",
		on_change=sincronizar_matriz
	)

	if st.button("💾 Guardar Calificaciones y Calcular Brechas", use_container_width=True, type="primary"):
		with st.spinner("Guardando calificaciones y calculando métricas RAG para toda la sección..."):
			for index, row in edited_df.iterrows():
				alumno_id = row["ID_Alumno"]
				puntajes_a_guardar = []
				acumulado_indicadores = {}

				for p in preguntas_ordenadas:
					num_preg = p["numero_pregunta"]
					col_name = f"{num_preg} (Max: {p['puntos_maximos']})"
					obtenido = float(row[col_name])
					maximo = float(p["puntos_maximos"])

					puntajes_a_guardar.append({"numero_pregunta": num_preg, "puntos_obtenidos": obtenido})

					# Acumular por indicador para este alumno
					inds_asociados = _parsear_indicadores_asociados(p["indicadores_asociados"])
					for ind in inds_asociados:
						if ind not in acumulado_indicadores:
							acumulado_indicadores[ind] = {"obtenido": 0.0, "maximo": 0.0}
						acumulado_indicadores[ind]["obtenido"] += obtenido
						acumulado_indicadores[ind]["maximo"] += maximo

				# Persistencia en BD
				guardar_puntajes_pregunta(alumno_id, evaluacion_id, puntajes_a_guardar)

				resumen_indicadores = []
				for ind, vals in acumulado_indicadores.items():
					logro = vals["obtenido"] / vals["maximo"] if vals["maximo"] > 0 else 0.0
					resumen_indicadores.append({
						"Indicador": ind,
						"Logro": round(logro, 4),
						"Observacion": f"{evaluacion['nombre']} | {vals['obtenido']:.2f}/{vals['maximo']:.2f}"
					})
				reemplazar_resultados_alumno(alumno_id, resumen_indicadores)

			# Limpiar el borrador temporal para forzar la recarga limpia desde la BD
			if clave_editor in st.session_state:
				del st.session_state[clave_editor]

			st.success("✅ Matriz de curso guardada. Diagnóstico colectivo actualizado.")
			st.rerun()

	resumen_indicadores: list[dict[str, Any]] = []

	if resumen_indicadores:
		st.markdown("### 2) Logro consolidado por indicador")
		st.dataframe(pd.DataFrame(resumen_indicadores), use_container_width=True)

	mapa_cohortes: dict[str, list[str]] = {}
	for alumno in alumnos:
		resultados = listar_resultados_alumno(int(alumno["id"]))
		if not resultados:
			continue
		nombre_alumno = str(alumno["nombre"]).strip()
		for fila in resultados:
			try:
				if float(fila.get("logro", 0)) >= 0.60:
					continue
			except (TypeError, ValueError):
				continue
			indicador = str(fila.get("indicador", "")).strip()
			if not indicador:
				continue
			claves = mapa_cohortes.setdefault(indicador, [])
			if nombre_alumno and nombre_alumno not in claves:
				claves.append(nombre_alumno)

	st.markdown("### 📊 Diagnóstico Colectivo y Planes de Reforzamiento")

	# 1. Analizar brechas por alumno y agrupar por patrón exacto
	brechas_por_alumno = {}
	for alumno in alumnos:
		alumno_id = int(alumno["id"])
		resultados = listar_resultados_alumno(alumno_id)
		brechas_locales = [str(r["indicador"]) for r in resultados if float(r.get("logro", 0)) < 0.60]
		if brechas_locales:
			brechas_por_alumno[alumno_id] = {
				"nombre": alumno["nombre"],
				"email": alumno.get("email", ""),
				"brechas": sorted(brechas_locales)
			}

	if not brechas_por_alumno:
		st.success("✅ Excelente rendimiento: Ningún alumno presenta brechas bajo el 60% en esta evaluación.")
		return

	# Agrupar por el mismo conjunto de brechas
	mapa_patrones = {}
	for a_id, data in brechas_por_alumno.items():
		patron = tuple(data["brechas"])
		if patron not in mapa_patrones:
			mapa_patrones[patron] = []
		mapa_patrones[patron].append(data)

	st.info("💡 El sistema ha agrupado automáticamente a los alumnos que fallaron exactamente en los mismos indicadores para optimizar los planes.")

	for patron, grupo_alumnos in mapa_patrones.items():
		nombres = [a["nombre"] for a in grupo_alumnos]
		st.markdown(f"**Grupo con brechas en: {', '.join(patron)}** ({len(nombres)} alumnos)")
		st.caption(", ".join(nombres))

	st.markdown("#### ⚙️ Motor de Generación Masiva")
	tau_generacion = st.slider("Umbral de similitud (Tau retrieval) para IA", 0.50, 0.95, 0.75, 0.01)

	# Estado de sesión para guardar los planes generados temporalmente
	if "planes_generados_zip" not in st.session_state:
		st.session_state.planes_generados_zip = None

	if st.button("🚀 Generar TODOS los Planes de Reforzamiento (Masivo)", use_container_width=True, type="primary"):
		progreso = st.progress(0, text="Iniciando motor de generación RAG...")
		total_grupos = len(mapa_patrones)

		planes_completados = []

		for i, (patron, grupo_alumnos) in enumerate(mapa_patrones.items()):
			nombres = [a["nombre"] for a in grupo_alumnos]
			emails = [a["email"] for a in grupo_alumnos if a["email"]]

			progreso.progress((i) / total_grupos, text=f"Generando plan {i+1} de {total_grupos} (Brechas: {', '.join(patron)})...")

			try:
				# Usa la función original, enviando la lista de nombres y forzando las brechas
				# Usa la función original, enviando la lista de nombres, forzando las brechas y la unidad activa
				salida = generar_plan_trabajo_cohorte(
					nombres,
					brechas_forzadas=list(patron),
					tau=tau_generacion,
					unidad_id=str(unidad["id"])
				)
				plan_texto = salida.get("plan_4_semanas", "")

				planes_completados.append({
					"patron": list(patron),
					"alumnos": nombres,
					"emails": emails,
					"contenido": plan_texto
				})
			except Exception as e:
				st.error(f"Error generando plan para grupo {', '.join(nombres)}: {e}")

		progreso.progress(1.0, text="¡Generación completada!")

		# Crear archivo ZIP en memoria
		zip_buffer = io.BytesIO()
		with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
			for idx, plan_data in enumerate(planes_completados):
				# Crear nombre de archivo basado en emails o nombres
				identificador = "_".join(plan_data["emails"]) if plan_data["emails"] else f"Grupo_{idx+1}"
				if not identificador.strip():
					identificador = f"Grupo_{idx+1}"

				nombre_archivo = f"Plan_Refuerzo_{identificador}.md"
				# Añadir metadatos al documento
				encabezado = f"Plan de Reforzamiento\nAlumnos: {', '.join(plan_data['alumnos'])}\nBrechas: {', '.join(plan_data['patron'])}\n\n"
				contenido_final = encabezado + plan_data["contenido"]

				zip_file.writestr(nombre_archivo, contenido_final)

		st.session_state.planes_generados_zip = zip_buffer.getvalue()
		st.session_state.vista_planes = planes_completados
		st.rerun()

	# Visualización en Acordeón (Expanders) si hay planes generados
	if st.session_state.get("planes_generados_zip"):
		st.success("✅ Planes generados exitosamente. Puedes revisarlos a continuación y descargar el paquete para envío.")

		st.download_button(
			label="📦 Descargar todos los planes (.ZIP para Email)",
			data=st.session_state.planes_generados_zip,
			file_name=f"Planes_Refuerzo_{evaluacion['nombre'].replace(' ', '_')}.zip",
			mime="application/zip",
			use_container_width=True
		)

		for idx, plan in enumerate(st.session_state.vista_planes):
			with st.expander(f"Plan {idx+1}: {', '.join(plan['alumnos'])} (Brechas: {', '.join(plan['patron'])})"):
				st.markdown(plan["contenido"])


def _vista_reportes() -> None:
	st.subheader("📊 Reporte de Jefatura")
	with st.container(border=True):
		st.info("Instrucciones de uso: revisa metricas de riesgo bajo 60% y utiliza las alertas para priorizar acompanamiento docente.")
	unidad = _unidad_activa()
	if not unidad:
		st.info("Activa una unidad para generar reportes.")
		return

	alumnos = listar_alumnos(anio=unidad["anio"], mencion=unidad["mencion"])
	total = len(alumnos)

	bajo_60_local = 0
	detalle_local: list[dict[str, Any]] = []
	conteo_brechas_indicador: dict[str, int] = {}
	logros_bajo_umbral: list[float] = []
	for alumno in alumnos:
		filas = listar_resultados_alumno(int(alumno["id"]))
		brechas = []
		for fila in filas:
			try:
				if float(fila.get("logro", 0)) < 0.60:
					indicador = str(fila.get("indicador", ""))
					brechas.append(indicador)
					conteo_brechas_indicador[indicador] = conteo_brechas_indicador.get(indicador, 0) + 1
					logros_bajo_umbral.append(float(fila.get("logro", 0)))
			except (TypeError, ValueError):
				continue
		if brechas:
			bajo_60_local += 1
			detalle_local.append(
				{
					"Alumno": alumno["nombre"],
					"RUT": alumno["rut"],
					"Brechas": ", ".join(brechas),
				}
			)

	bajo_60_excel = 0
	try:
		for nombre in obtener_alumnos():
			if obtener_brechas_alumno(nombre).get("indicadores_brecha"):
				bajo_60_excel += 1
	except Exception:
		pass

	porcentaje_riesgo_local = (bajo_60_local / total * 100.0) if total else 0.0
	logro_promedio_bajo = (sum(logros_bajo_umbral) / len(logros_bajo_umbral) * 100.0) if logros_bajo_umbral else 0.0
	delta_logro = f"-{100.0 - logro_promedio_bajo:.1f}% de logro promedio" if logros_bajo_umbral else "Sin rezago critico"

	if bajo_60_local > 0:
		st.error(
			f"⚠️ Alerta: Se ha detectado {bajo_60_local} estudiante"
			f"{'s' if bajo_60_local != 1 else ''} con brechas por debajo del umbral mínimo del 60%. "
			"Requiere acompañamiento pedagógico."
		)

	c1, c2, c3 = st.columns(3)
	with c1:
		with st.container(border=True):
			st.metric("Alumnos en unidad activa", total)
	with c2:
		with st.container(border=True):
			st.metric(
				"Riesgo Local",
				bajo_60_local,
				delta=delta_logro,
			)
			if bajo_60_local > 0:
				st.warning(f"{porcentaje_riesgo_local:.1f}% del curso presenta alertas locales.")
	with c3:
		with st.container(border=True):
			st.metric("Riesgo Excel", bajo_60_excel)

	with st.container(border=True):
		st.markdown("### 📊 Diagnóstico Colectivo de Indicadores")
		if not conteo_brechas_indicador:
			st.info("No se detectan brechas colectivas bajo 60% en la unidad activa.")
		else:
			df_brechas = pd.DataFrame(
				[
					{"Indicador": indicador, "Alumnos con brecha": cantidad}
					for indicador, cantidad in conteo_brechas_indicador.items()
				]
			)
			df_brechas = df_brechas.sort_values(by="Alumnos con brecha", ascending=False)
			try:
				import plotly.express as px

				fig = px.bar(
					df_brechas,
					x="Alumnos con brecha",
					y="Indicador",
					orientation="h",
					text="Alumnos con brecha",
				)
				fig.update_layout(height=360, margin=dict(l=20, r=20, t=20, b=20))
				st.plotly_chart(fig, use_container_width=True)
			except Exception:
				st.bar_chart(df_brechas.set_index("Indicador"), use_container_width=True)

	with st.container(border=True):
		with st.expander("🔍 Detalle Nominal de Alumnos en Riesgo", expanded=False):
			if not detalle_local:
				st.success("Sin alertas locales para la unidad activa.")
			else:
				st.dataframe(pd.DataFrame(detalle_local), use_container_width=True)


def main() -> None:
	_inject_styles()
	_init_state()
	_bloquear_lector_navegador()  # <- Nueva línea agregada

	if not st.session_state.is_logged:
		_render_login()
		return

	_render_top_context()
	modulo = _menu_sidebar()

	if st.session_state.usuario_rol == "admin":
		if modulo.startswith("1"):
			_vista_admin_matricula()
		else:
			_vista_admin_cuentas()
		return

	if modulo.startswith("1"):
		_vista_unidades()
	else:
		# Bloqueo estricto si no hay unidad activa en módulos 2, 3, 4 y 5
		if not st.session_state.get("unidad_activa_id"):
			st.error("⚠️ **ACCESO DENEGADO:** Debes crear o seleccionar una 'Unidad Activa' en el módulo 'Gestión de Unidades' antes de poder acceder a las demás funciones del sistema.")
			return

		if modulo.startswith("2"):
			_vista_bibliografia()
		elif modulo.startswith("3"):
			_vista_diseno_evaluacion()
		elif modulo.startswith("4"):
			_vista_trazabilidad()
		elif modulo.startswith("5"):
			_vista_reportes()


if __name__ == "__main__":
	main()
