from __future__ import annotations

from dataclasses import dataclass
from hashlib import sha256
from pathlib import Path
from typing import Iterable
import math
import os

# Bloqueo de telemetría de ChromaDB
os.environ["ANONYMIZED_TELEMETRY"] = "False"

import docx2txt
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_core.documents import Document
from langchain_core.embeddings import Embeddings
from langchain_community.vectorstores import Chroma
from langchain_community.embeddings import HuggingFaceEmbeddings

ROOT_DIR = Path(__file__).resolve().parent.parent
DOCTRINA_DIR = ROOT_DIR / "data" / "doctrina"
VECTOR_STORE_DIR = ROOT_DIR / "data" / "vector_store"
COLLECTION_NAME = "acapomil_doctrina"
TAU_SIMILITUD = 0.60  # Umbral táctico ajustado para doctrina en español


class HashEmbeddings(Embeddings):
	"""Embeddings deterministas de respaldo si HuggingFace no está disponible."""
	def __init__(self, dim: int = 256) -> None:
		self.dim = dim

	def _embed_text(self, text: str) -> list[float]:
		vector = [0.0] * self.dim
		if not text:
			return vector
		tokens = text.lower().split()
		for token in tokens:
			digest = sha256(token.encode("utf-8")).digest()
			idx = int.from_bytes(digest[:4], "big") % self.dim
			sign = 1.0 if digest[4] % 2 == 0 else -1.0
			vector[idx] += sign
		norm = math.sqrt(sum(v * v for v in vector))
		if norm == 0:
			return vector
		return [v / norm for v in vector]

	def embed_documents(self, texts: list[str]) -> list[list[float]]:
		return [self._embed_text(text) for text in texts]

	def embed_query(self, text: str) -> list[float]:
		return self._embed_text(text)


@dataclass(frozen=True)
class ResultadoRecuperacion:
	fragmento: str
	similitud_coseno: float
	metadata: dict


def get_embeddings() -> Embeddings:
	"""Utiliza embeddings locales multilingües optimizados en CPU."""
	try:
		return HuggingFaceEmbeddings(
			# Modelo optimizado para comprender semántica en Español
			model_name="sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2",
			model_kwargs={"device": "cpu"},
			encode_kwargs={"batch_size": 64}
		)
	except Exception as e:
		print(f"[RAG] Error cargando HuggingFaceEmbeddings: {e}")
		return HashEmbeddings()


def _extraer_texto_pdf(path_archivo: Path) -> str:
	from pypdf import PdfReader
	try:
		reader = PdfReader(path_archivo)
		texto = ""
		for pagina in reader.pages:
			texto += (pagina.extract_text() or "") + "\n"
		return texto.strip()
	except Exception as e:
		print(f"Error al leer PDF {path_archivo}: {e}")
		return ""


def _extraer_texto_pptx(path_archivo: Path) -> str:
	from pptx import Presentation
	try:
		prs = Presentation(path_archivo)
		texto = ""
		for slide in prs.slides:
			for shape in slide.shapes:
				if hasattr(shape, "text") and shape.text.strip():
					texto += shape.text.strip() + "\n"
			if slide.notes_slide and slide.notes_slide.notes_text_frame:
				texto += slide.notes_slide.notes_text_frame.text.strip() + "\n"
		return texto.strip()
	except Exception as e:
		print(f"Error al leer PPTX {path_archivo}: {e}")
		return ""


def cargar_documentos_docx(origen: Path = DOCTRINA_DIR) -> list[Document]:
	if not origen.exists():
		raise FileNotFoundError(f"No existe la carpeta de doctrina: {origen}")
	documentos: list[Document] = []
	extensiones_permitidas = {".docx", ".pdf", ".pptx", ".ppt"}
	for ruta in sorted(p for p in origen.iterdir() if p.is_file() and p.suffix.lower() in extensiones_permitidas):
		texto = ""
		sufijo = ruta.suffix.lower()
		if sufijo == ".docx":
			texto = docx2txt.process(str(ruta))
		elif sufijo == ".pdf":
			texto = _extraer_texto_pdf(ruta)
		elif sufijo in {".pptx", ".ppt"}:
			texto = _extraer_texto_pptx(ruta)
		if not texto or not texto.strip():
			continue
		documentos.append(
			Document(
				page_content=texto,
				metadata={"source": str(ruta), "filename": ruta.name},
			)
		)
	if not documentos:
		raise ValueError(f"No se encontraron documentos con contenido en {origen}")
	return documentos


def crear_segmentador() -> RecursiveCharacterTextSplitter:
	return RecursiveCharacterTextSplitter(
		chunk_size=2500,
		chunk_overlap=250,
		separators=["\n\n", "\n", ". ", " ", ""],
	)


def segmentar_documentos(documentos: Iterable[Document]) -> list[Document]:
	splitter = crear_segmentador()
	return splitter.split_documents(list(documentos))


def construir_vector_store(
	documentos: list[Document],
	persist_directory: Path = VECTOR_STORE_DIR,
	collection_name: str = COLLECTION_NAME,
	progress_callback=None,
) -> Chroma:
	persist_directory.mkdir(parents=True, exist_ok=True)
	embeddings = get_embeddings()
	chunks = segmentar_documentos(documentos)
	total_chunks = len(chunks)

	msg_init = f"Iniciando vectorización masiva de {total_chunks} fragmentos..."
	print(f"\n[RAG] {msg_init}")
	if progress_callback:
		progress_callback(0.0, msg_init)

	# Fragmentado en lotes para no exceder el límite de variables SQLite de Chroma (max batch size 5461)
	TAMANO_LOTE = 500
	vector_store = None
	total_lotes = math.ceil(total_chunks / TAMANO_LOTE)

	for i in range(0, total_chunks, TAMANO_LOTE):
		lote = chunks[i : i + TAMANO_LOTE]
		lote_actual = (i // TAMANO_LOTE) + 1

		if progress_callback:
			progress_callback((lote_actual - 1) / total_lotes, f"Procesando lote {lote_actual} de {total_lotes}...")

		if vector_store is None:
			vector_store = Chroma.from_documents(
				documents=lote,
				embedding=embeddings,
				collection_name=collection_name,
				persist_directory=str(persist_directory),
			)
		else:
			vector_store.add_documents(documents=lote)

	msg_fin = "🎯 Vectorización completada con éxito."
	print(f"[RAG] {msg_fin}\n")
	if progress_callback:
		progress_callback(1.0, msg_fin)

	return vector_store


def cargar_vector_store(
	persist_directory: Path = VECTOR_STORE_DIR,
	collection_name: str = COLLECTION_NAME,
) -> Chroma:
	if not persist_directory.exists():
		raise FileNotFoundError(f"No existe la carpeta del vector store: {persist_directory}.")
	return Chroma(
		collection_name=collection_name,
		embedding_function=get_embeddings(),
		persist_directory=str(persist_directory),
	)


def _similitud_coseno(vec_a: list[float], vec_b: list[float]) -> float:
	if not vec_a or not vec_b or len(vec_a) != len(vec_b):
		return 0.0
	producto = sum(a * b for a, b in zip(vec_a, vec_b))
	norma_a = math.sqrt(sum(a * a for a in vec_a))
	norma_b = math.sqrt(sum(b * b for b in vec_b))
	if norma_a == 0 or norma_b == 0:
		return 0.0
	return producto / (norma_a * norma_b)


def obtener_conteo_fragmentos_unidad(unidad_id: str) -> int:
	"""Consulta cuántos fragmentos vectoriales existen persistidos para la unidad."""
	vector_dir = VECTOR_STORE_DIR / str(unidad_id)
	col_name = f"acapomil_{unidad_id}"
	if not vector_dir.exists():
		return 0
	try:
		store = cargar_vector_store(persist_directory=vector_dir, collection_name=col_name)
		datos = store.get()
		return len(datos.get("ids", []))
	except Exception:
		return 0


def recuperar_fragmentos_relevantes(
	pregunta: str,
	tau: float = TAU_SIMILITUD,
	k: int = 4,
	vector_store: Chroma | None = None,
) -> list[ResultadoRecuperacion]:
	if not pregunta or not pregunta.strip():
		raise ValueError("La pregunta no puede estar vacía.")
	store = vector_store or cargar_vector_store()
	embeddings = get_embeddings()
	candidatos = store.similarity_search(query=pregunta, k=k)
	if not candidatos:
		return []
	emb_pregunta = embeddings.embed_query(pregunta)
	resultados: list[ResultadoRecuperacion] = []
	for doc in candidatos:
		emb_doc = embeddings.embed_query(doc.page_content)
		similitud = _similitud_coseno(emb_pregunta, emb_doc)
		if similitud >= tau:
			resultados.append(
				ResultadoRecuperacion(
					fragmento=doc.page_content,
					similitud_coseno=similitud,
					metadata=doc.metadata,
				)
			)
	resultados.sort(key=lambda item: item.similitud_coseno, reverse=True)
	return resultados


def ingest_arreglo_doctrina() -> Chroma:
	documentos = cargar_documentos_docx()
	return construir_vector_store(documentos=documentos)
